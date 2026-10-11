"""Read the recalculated audit workbook, draw the report figures and export report data.

Usage: python3 make_report_assets.py <workbook.xlsx> <figures_dir> <report_data.json>
Every number used in the Word report comes from the workbook through this script.
"""
import json
import os
import sys
from collections import Counter

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

from audit_inputs import AUDIT, FLIGHTS, INPUTS, MONTHS, SOURCES  # noqa: E402

WB, FIG, OUT = sys.argv[1:4]
os.makedirs(FIG, exist_ok=True)
wb = load_workbook(WB, data_only=True)

# ---------------- palette (dataviz reference palette, light mode, fixed order)
SER = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"]
INK, INK2, GRID, SURF = "#0b0b0b", "#52514e", "#e2e1dc", "#ffffff"
POS, NEG = "#e34948", "#2a78d6"  # diverging poles: red = more emissions, blue = fewer
plt.rcParams.update({
    "font.family": "Liberation Sans", "font.size": 9, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
    "xtick.color": INK2, "ytick.color": INK2, "axes.titlesize": 10, "axes.titleweight": "bold",
    "axes.titlecolor": INK, "legend.frameon": False, "figure.dpi": 200, "savefig.dpi": 200,
    "axes.spines.top": False, "axes.spines.right": False,
})


def style(ax, axis="y"):
    ax.set_axisbelow(True)
    ax.grid(axis=axis, color=GRID, linewidth=0.6)
    ax.tick_params(length=0)
    for sp in ("left", "bottom"):
        ax.spines[sp].set_color(GRID)


def save(fig, name):
    path = os.path.join(FIG, name)
    fig.savefig(path, bbox_inches="tight", facecolor=SURF)
    plt.close(fig)
    return name


def v(sheet, ref):
    return wb[sheet][ref].value


def fmt(x, d=0):
    return f"{x:,.{d}f}"


D = {"audit": AUDIT}

# ---------------- headline numbers
k = {}
k["ac_td"] = v("AC_Baseline", "K9")
k["ac_bu"] = v("AC_Baseline", "K10")
k["ac_gap"] = v("AC_Baseline", "K11")
k["ac_stated"] = v("AC_Baseline", "K12")
k["ac_implied_tariff"] = v("AC_Baseline", "K13")
k["ac_avg_kw_run"] = v("AC_Baseline", "K14")
k["ac_eflh"] = v("AC_Baseline", "K15")
k["ac_hot"] = v("AC_Baseline", "K16")
k["ac_hot_share"] = v("AC_Baseline", "K17")
k["ac_cool"] = v("AC_Baseline", "K18")
k["ac_spend"] = v("AC_Baseline", "K19")
k["ac_hours"] = v("AC_Baseline", "K8")
k["ac_label_kw"] = v("AC_Baseline", "K6")
k["ac_cooling_th"] = v("AC_Baseline", "K20")
k["hw_kwh"] = v("HotWater", "G18")
k["hw_avg_dt"] = v("HotWater", "D18")
k["hw_peak"] = v("HotWater", "H18")
k["hw_headroom"] = v("HotWater", "K8")
k["hw_jan_jul"] = v("HotWater", "K12")
k["hw_pp_day"] = v("HotWater", "F18")
k["hw_l_pp_day"] = v("HotWater", "E6")
k["hh_kwh"] = v("Inventory", "J23")
k["pp_kwh"] = v("Inventory", "K23")
k["room_kwh"] = v("Inventory", "D27")
k["wr_kwh"] = v("Inventory", "D28")
k["kit_kwh"] = v("Inventory", "D29")
k["lau_kwh"] = v("Inventory", "D30")
k["shared_kwh"] = v("Inventory", "D32")
k["shared_pct"] = v("Inventory", "E32")
k["room_pct"] = v("Inventory", "E27")
k["room_eui_kwh"] = v("Breakdown", "B30")
k["room_eui_mj"] = v("Breakdown", "B31")
k["bm_eui_kwh"] = v("Breakdown", "D30")
k["bm_eui_mj"] = v("Breakdown", "D31")
k["eui_ratio"] = v("Breakdown", "F30")
k["bm_pp_kwh"] = v("Breakdown", "D32")
k["pp_ratio"] = v("Breakdown", "F32")
k["lpd"] = v("Breakdown", "B34")
k["bec_lpd"] = v("Breakdown", "D34")
k["floor_pp"] = v("Breakdown", "B35")
k["bm_floor_pp"] = v("Breakdown", "D35")
k["water_lpd"] = v("Breakdown", "B36")
k["bm_water"] = v("Breakdown", "D36")
k["ac_share_room"] = v("Breakdown", "B37")
k["hw_share"] = v("Breakdown", "B38")
k["s1"] = v("GHG_Inventory", "B53")
k["s2_clp"] = v("GHG_Inventory", "B54")
k["s2_def"] = v("GHG_Inventory", "C54")
k["s3"] = v("GHG_Inventory", "B55")
k["tot_clp"] = v("GHG_Inventory", "B56")
k["tot_def"] = v("GHG_Inventory", "C56")
k["tot_clp_pp"] = v("GHG_Inventory", "F56")
k["tot_def_pp"] = v("GHG_Inventory", "G56")
k["s2_share"] = v("GHG_Inventory", "E54")
k["room_intensity"] = v("GHG_Inventory", "B57")
k["fresh_m3"] = v("GHG_Inventory", "B41")
k["salt_m3"] = v("GHG_Inventory", "B42")
k["paper_kg"] = v("GHG_Inventory", "B36")
k["life_kg"] = v("Lifestyle_Scope3", "F10")
k["mtr_kg"] = v("Lifestyle_Scope3", "F6") + v("Lifestyle_Scope3", "F7")
k["mtr_km"] = v("Lifestyle_Scope3", "B6") + v("Lifestyle_Scope3", "B7")
k["fl_kg"] = v("Lifestyle_Scope3", "F18")
k["fl_km"] = v("Lifestyle_Scope3", "D18")
k["diet_kg"] = v("Lifestyle_Scope3", "F9")
k["personal_kg"] = v("Lifestyle_Scope3", "B26")
k["home_pp_kg"] = v("Lifestyle_Scope3", "B22")
k["home_share_personal"] = v("Lifestyle_Scope3", "C22")
k["hk_pc_kg"] = v("Lifestyle_Scope3", "B27")
k["personal_ratio"] = v("Lifestyle_Scope3", "B28")
inputs = {r[0]: r[3] for r in INPUTS}
k.update({f"in_{n}": val for n, val in inputs.items()})
D["k"] = k

# ---------------- tables
inv = []
for r in range(7, 23):
    inv.append(dict(id=v("Inventory", f"A{r}"), space=v("Inventory", f"B{r}"), item=v("Inventory", f"C{r}"),
                    model=v("Inventory", f"D{r}"), power=v("Inventory", f"E{r}"), qty=v("Inventory", f"F{r}"),
                    use=v("Inventory", f"G{r}"), total=v("Inventory", f"H{r}"), share=v("Inventory", f"I{r}"),
                    hh=v("Inventory", f"J{r}"), pp=v("Inventory", f"K{r}"), enduse=v("Inventory", f"L{r}"),
                    eac=v("Inventory", f"M{r}"), quality=v("Inventory", f"N{r}"), photo=v("Inventory", f"O{r}")))
D["inventory"] = inv
monthly = []
for i in range(12):
    r = 6 + i
    row = {c: v("Monthly", f"{c}{r}") for c in "ABCDEFGHIJKLMNOPQRST"}
    monthly.append(row)
D["monthly"] = monthly
D["monthly_total"] = {c: v("Monthly", f"{c}18") for c in "BDEFGHIJKLMNOPQRST"}
D["ac_months"] = [dict(m=v("AC_Baseline", f"A{6 + i}"), season=v("AC_Baseline", f"C{6 + i}"),
                       spend=v("AC_Baseline", f"D{6 + i}"), td=v("AC_Baseline", f"E{6 + i}"),
                       hours=v("AC_Baseline", f"F{6 + i}"), bu=v("AC_Baseline", f"G{6 + i}")) for i in range(12)]
D["hw_months"] = [dict(m=v("HotWater", f"A{6 + i}"), tin=v("HotWater", f"C{6 + i}"), dt=v("HotWater", f"D{6 + i}"),
                       kwh_pp_day=v("HotWater", f"F{6 + i}"), kwh=v("HotWater", f"G{6 + i}"),
                       peak=v("HotWater", f"H{6 + i}")) for i in range(12)]
D["enduse"] = [dict(name=v("Breakdown", f"A{r}"), kwh=v("Breakdown", f"B{r}"), pct=v("Breakdown", f"C{r}"),
                    pp=v("Breakdown", f"D{r}"), kg=v("Breakdown", f"E{r}"), kg_def=v("Breakdown", f"F{r}"))
               for r in range(6, 16)]
D["eac"] = [dict(name=v("Breakdown", f"A{r}"), room=v("Breakdown", f"B{r}"), room_pct=v("Breakdown", f"C{r}"),
                 hh=v("Breakdown", f"D{r}"), hh_pct=v("Breakdown", f"E{r}"), room_mj=v("Breakdown", f"F{r}"))
            for r in range(21, 25)]
D["eac_total"] = dict(room=v("Breakdown", "B25"), hh=v("Breakdown", "D25"), room_mj=v("Breakdown", "F25"))
D["scope1"] = [dict(eq=v("GHG_Inventory", f"A{r}"), ref=v("GHG_Inventory", f"B{r}"), charge=v("GHG_Inventory", f"C{r}"),
                    leak=v("GHG_Inventory", f"D{r}"), gwp=v("GHG_Inventory", f"E{r}"), share=v("GHG_Inventory", f"F{r}"),
                    kg=v("GHG_Inventory", f"G{r}")) for r in (17, 18)]
D["scope2"] = [dict(name=v("GHG_Inventory", f"A{r}"), kwh=v("GHG_Inventory", f"B{r}"), clp=v("GHG_Inventory", f"E{r}"),
                    dflt=v("GHG_Inventory", f"F{r}")) for r in range(23, 28)]
D["scope3"] = [dict(name=v("GHG_Inventory", f"A{r}"), q=v("GHG_Inventory", f"B{r}"), unit=v("GHG_Inventory", f"C{r}"),
                    ef=v("GHG_Inventory", f"D{r}"), efu=v("GHG_Inventory", f"E{r}"), kg=v("GHG_Inventory", f"F{r}"),
                    basis=v("GHG_Inventory", f"G{r}")) for r in range(45, 49)]
D["water"] = [dict(name=v("GHG_Inventory", f"A{r}"), q=v("GHG_Inventory", f"B{r}"), unit=v("GHG_Inventory", f"C{r}"))
              for r in (32, 33, 34, 35, 40, 41, 42)]
D["ghg_summary"] = [dict(name=v("GHG_Inventory", f"A{r}"), clp=v("GHG_Inventory", f"B{r}"), dflt=v("GHG_Inventory", f"C{r}"),
                         pct=v("GHG_Inventory", f"E{r}"), pp=v("GHG_Inventory", f"F{r}"), pp_def=v("GHG_Inventory", f"G{r}"))
                    for r in range(53, 57)]
D["life"] = [dict(name=v("Lifestyle_Scope3", f"A{r}"), q=v("Lifestyle_Scope3", f"B{r}"), unit=v("Lifestyle_Scope3", f"C{r}"),
                  ef=v("Lifestyle_Scope3", f"D{r}"), efu=v("Lifestyle_Scope3", f"E{r}"), kg=v("Lifestyle_Scope3", f"F{r}"),
                  basis=v("Lifestyle_Scope3", f"G{r}")) for r in range(6, 11)]
D["flights"] = [dict(leg=v("Lifestyle_Scope3", f"A{r}"), gcd=v("Lifestyle_Scope3", f"B{r}"), haul=v("Lifestyle_Scope3", f"C{r}"),
                     km=v("Lifestyle_Scope3", f"D{r}"), ef=v("Lifestyle_Scope3", f"E{r}"), kg=v("Lifestyle_Scope3", f"F{r}"))
                for r in range(14, 18)]
D["personal"] = [dict(name=v("Lifestyle_Scope3", f"A{r}"), kg=v("Lifestyle_Scope3", f"B{r}"), pct=v("Lifestyle_Scope3", f"C{r}"))
                 for r in range(22, 27)]
emos = []
for r in range(6, 24):
    emos.append({key: v("EMO_Register", f"{col}{r}") for key, col in zip(
        ["ref", "cat", "type", "resp", "desc", "basis", "fac_kwh", "hh_kwh", "hh_hkd", "cost", "payback", "life",
         "fac_kg", "hh_kg", "rec"], "ABCDEFGHIJKLMNO")})
D["emos"] = emos
D["emo_summary"] = [dict(cat=v("EMO_Register", f"A{r}"), n=v("EMO_Register", f"B{r}"), hh_kwh=v("EMO_Register", f"C{r}"),
                         pct=v("EMO_Register", f"D{r}"), kg=v("EMO_Register", f"E{r}"), fac_kwh=v("EMO_Register", f"F{r}"))
                    for r in range(27, 31)]
D["emo_types"] = [dict(type=v("EMO_Register", f"A{r}"), n=v("EMO_Register", f"B{r}")) for r in range(34, 39)]
D["sensitivity"] = [dict(i=v("Sensitivity", f"A{r}"), name=v("Sensitivity", f"B{r}"), val=v("Sensitivity", f"C{r}"),
                         ac=v("Sensitivity", f"D{r}"), hh=v("Sensitivity", f"E{r}"), kg=v("Sensitivity", f"F{r}"),
                         chg=v("Sensitivity", f"G{r}")) for r in range(6, 20)]
D["inputs"] = [dict(name=r[0], section=r[1], label=r[2], value=r[3], unit=r[4], quality=r[5], source=r[6], note=r[7])
               for r in INPUTS]
D["quality_counts"] = dict(Counter(r[5] for r in INPUTS))
D["sources"] = [dict(id=s[0], short=s[1], cit=s[2], url=s[3]) for s in SOURCES]
D["entry_guide"] = [dict(sheet=v("EE-EAes_Entry_Guide", f"A{r}"), cell=v("EE-EAes_Entry_Guide", f"B{r}"),
                         field=v("EE-EAes_Entry_Guide", f"C{r}"), value=v("EE-EAes_Entry_Guide", f"D{r}"))
                    for r in range(5, wb["EE-EAes_Entry_Guide"].max_row + 1) if v("EE-EAes_Entry_Guide", f"A{r}")]
D["photo_log"] = [dict(ref=v("Photo_Log", f"A{r}"), file=v("Photo_Log", f"B{r}"), what=v("Photo_Log", f"C{r}"),
                       data=v("Photo_Log", f"D{r}")) for r in range(5, wb["Photo_Log"].max_row + 1) if v("Photo_Log", f"A{r}")]
D["flights_in"] = FLIGHTS
D["months"] = MONTHS

# sanity checks against the inventory
assert abs(sum(e["kwh"] for e in D["enduse"]) - k["hh_kwh"]) < 1e-6
assert abs(D["monthly_total"]["N"] - k["hh_kwh"]) < 1e-6
assert abs(k["s1"] + k["s2_clp"] + k["s3"] - k["tot_clp"]) < 1e-6

# ---------------- figures
figs = {}

# Fig: audit boundary schematic
fig, ax = plt.subplots(figsize=(7.2, 4.1))
ax.set_xlim(0, 100)
ax.set_ylim(-9, 58)
ax.axis("off")


def box(x, y, w, h, text, fc, ec, lw=1.2, fs=8.2, weight="normal", color=INK, ls="-"):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.6", fc=fc, ec=ec, lw=lw, ls=ls))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, color=color, weight=weight, wrap=True)


box(2, 4, 62, 50, "", "#f4f8fd", SER[0], lw=1.6)
ax.text(33, 50.5, "Household boundary: Room 1003B (2 occupants)", ha="center", fontsize=9.2, weight="bold", color=INK)
box(5, 33, 26, 13, "Room 1003B (15 m²)\n100% of AC, lighting,\nfan and plug loads", "#ffffff", SER[0])
box(35, 33, 26, 13, "Shared washroom (4 users)\nown showers +\n2/4 of fan and lights", "#ffffff", SER[1])
box(5, 15, 26, 13, "Shared kitchen (32 users)\nown cooking + 2/32 of\nfridges, dispenser, AC, lights", "#ffffff", SER[2], fs=7.8)
box(35, 15, 26, 13, "Hall laundry\nown wash and\ndryer loads", "#ffffff", SER[3])
ax.text(33, 8, "Excluded: corridors, lifts, lobbies, central plant, other residents' use",
        ha="center", fontsize=7.6, color=INK2, style="italic")
box(70, 40, 28, 12, "Scope 2\nCLP electricity\n(0.34 and 0.7 kg CO2e/kWh)", "#ffffff", INK2, fs=7.8)
box(70, 25, 28, 12, "Scope 1\nrefrigerant leakage\n(R32, R-410A)", "#ffffff", INK2, fs=7.8)
box(70, 10, 28, 12, "Scope 3\nwater, sewage, paper", "#ffffff", INK2, fs=7.8)
box(70, -6, 28, 11, "Reported separately:\ncommuting, flights, diet", "#ffffff", INK2, fs=7.6, ls="--", color=INK2)
for y in (46, 31, 16):
    ax.annotate("", xy=(69.5, y), xytext=(64.6, y), arrowprops=dict(arrowstyle="-|>", color=INK2, lw=1))
figs["boundary"] = save(fig, "fig01_boundary.png")

# Fig: AC monthly top-down vs bottom-up
fig, ax = plt.subplots(figsize=(7.2, 3.1))
x = range(12)
w = 0.36
td = [m["td"] for m in D["ac_months"]]
bu = [m["bu"] for m in D["ac_months"]]
ax.bar([i - w / 2 - 0.01 for i in x], td, w, color=SER[0], label="Top-down: Octopus spend ÷ HK$1.406/kWh", edgecolor=SURF, linewidth=1)
ax.bar([i + w / 2 + 0.01 for i in x], bu, w, color=SER[1], label="Bottom-up: label 0.302 kW × reported hours", edgecolor=SURF, linewidth=1)
ax.set_xticks(list(x), MONTHS)
ax.set_ylabel("kWh per month")
style(ax)
ax.legend(loc="upper left", ncol=2, fontsize=8, bbox_to_anchor=(0, 1.13))
ax.text(5.5, max(bu) * 1.02, f"Year: top-down {fmt(k['ac_td'])} kWh, bottom-up {fmt(k['ac_bu'])} kWh ({k['ac_gap']:+.1%})",
        ha="center", va="bottom", fontsize=8, color=INK)
ax.set_ylim(0, max(bu) * 1.15)
figs["ac"] = save(fig, "fig02_ac_monthly.png")

# Fig: hot water monthly
fig, ax = plt.subplots(figsize=(7.2, 2.9))
hw = [m["kwh"] for m in D["hw_months"]]
ax.bar(list(x), hw, 0.5, color=SER[0])
ax.set_xticks(list(x), [f"{m}\n{h['tin']:.1f}°C" for m, h in zip(MONTHS, D["hw_months"])])
ax.set_ylabel("kWh per month (household)")
style(ax)
for i in (0, 6):
    ax.text(i, hw[i] + 1, f"{hw[i]:.0f}", ha="center", va="bottom", fontsize=8, color=INK)
ax.set_ylim(0, max(hw) * 1.18)
ax.text(11.4, max(hw) * 1.12, "Second line = cold-water inlet (HKO monthly mean)", ha="right", fontsize=7.6, color=INK2)
figs["hw"] = save(fig, "fig03_hotwater_monthly.png")

# Fig: end-use horizontal bars
eu = sorted(D["enduse"], key=lambda e: e["kwh"])
fig, ax = plt.subplots(figsize=(7.2, 3.4))
ax.barh([e["name"] for e in eu], [e["kwh"] for e in eu], 0.55, color=SER[0])
for i, e in enumerate(eu):
    ax.text(e["kwh"] + 8, i, f"{e['kwh']:,.0f} kWh ({e['pct']:.1%})", va="center", fontsize=8, color=INK)
ax.set_xlabel("kWh per year (household)")
style(ax, "x")
ax.set_xlim(0, max(e["kwh"] for e in eu) * 1.3)
figs["enduse"] = save(fig, "fig04_enduse.png")

# Fig: monthly stacked (6 groups)
groups = [
    ("Room AC", ["D"]),
    ("Hot water (showers)", ["H"]),
    ("Cooking", ["L"]),
    ("Laptops, phones and fan", ["E", "G"]),
    ("Room lighting", ["F"]),
    ("Shared washroom, kitchen and laundry", ["I", "J", "K", "M"]),
]
fig, ax = plt.subplots(figsize=(7.2, 3.6))
bottom = [0.0] * 12
for gi, (lab, cols) in enumerate(groups):
    vals = [sum(m[c] for c in cols) for m in monthly]
    ax.bar(list(x), vals, 0.56, bottom=bottom, color=SER[gi], label=lab, edgecolor=SURF, linewidth=1.2)
    bottom = [b + vv for b, vv in zip(bottom, vals)]
for i in (0, 9):
    ax.text(i, bottom[i] + 3, f"{bottom[i]:.0f}", ha="center", va="bottom", fontsize=8, color=INK)
ax.set_xticks(list(x), MONTHS)
ax.set_ylabel("kWh per month (household)")
style(ax)
ax.set_ylim(0, max(bottom) * 1.13)
ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=3, fontsize=7.8)
figs["monthly"] = save(fig, "fig05_monthly_stacked.png")

# Fig: monthly room EUI
fig, ax = plt.subplots(figsize=(7.2, 2.6))
eui = [m["Q"] for m in monthly]
ax.bar(list(x), eui, 0.5, color=SER[0])
for i in (0, 4):
    ax.text(i, eui[i] + 0.4, f"{eui[i]:.1f}", ha="center", va="bottom", fontsize=8, color=INK)
ax.set_xticks(list(x), MONTHS)
ax.set_ylabel("MJ/m² per month")
style(ax)
ax.set_ylim(0, max(eui) * 1.2)
figs["eui"] = save(fig, "fig06_monthly_eui.png")

# Fig: EAC categories 100% stacked (room vs household)
cats = [e for e in D["eac"] if e["name"] != "Lift & escalator"]
fig, ax = plt.subplots(figsize=(7.2, 1.9))
rows = [("Household (incl. shared services)", "hh_pct"), ("Room 1003B (15 m²)", "room_pct")]
for ri, (lab, key) in enumerate(rows):
    left = 0
    for ci, c in enumerate(cats):
        val = c[key]
        ax.barh(ri, val, 0.5, left=left, color=SER[ci], edgecolor=SURF, linewidth=1.5, label=c["name"] if ri == 0 else None)
        if val > 0.07:
            ax.text(left + val / 2, ri, f"{val:.0%}", ha="center", va="center", fontsize=8.5,
                    color="#ffffff" if ci == 0 else INK)
        left += val
ax.set_yticks([0, 1], [r[0] for r in rows])
ax.set_xlim(0, 1)
ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
style(ax, "x")
ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.25), ncol=3, fontsize=8)
figs["eac"] = save(fig, "fig07_eac_categories.png")

# Fig: Scope 2 by facility, CLP vs default
s2 = [s for s in D["scope2"] if s["name"] != "Scope 2 total"][::-1]
fig, ax = plt.subplots(figsize=(7.2, 2.6))
y = range(len(s2))
h = 0.34
ax.barh([i + h / 2 + 0.01 for i in y], [s["clp"] for s in s2], h, color=SER[0], label="CLP-specific factor 0.34 kg/kWh (2025)")
ax.barh([i - h / 2 - 0.01 for i in y], [s["dflt"] for s in s2], h, color=SER[1], label="Territory-wide default 0.7 kg/kWh")
for i, s in enumerate(s2):
    ax.text(s["dflt"] + 6, i - h / 2, f"{s['dflt']:,.0f}", va="center", fontsize=7.8, color=INK)
    ax.text(s["clp"] + 6, i + h / 2, f"{s['clp']:,.0f}", va="center", fontsize=7.8, color=INK)
ax.set_yticks(list(y), [s["name"].split(" (")[0] for s in s2])
ax.set_xlabel("kg CO2e per year")
style(ax, "x")
ax.set_xlim(0, max(s["dflt"] for s in s2) * 1.15)
ax.legend(loc="lower right", fontsize=7.8)
figs["scope2"] = save(fig, "fig08_scope2_facility.png")

# Fig: personal footprint
pf = sorted([p for p in D["personal"] if p["name"] != "Personal total"], key=lambda p: p["kg"])
labels = {"Home energy & water (household inventory ÷ 2)": "Home energy & water (this audit)"}
fig, ax = plt.subplots(figsize=(7.2, 2.3))
ax.barh([labels.get(p["name"], p["name"]) for p in pf], [p["kg"] for p in pf], 0.5, color=SER[0])
for i, p in enumerate(pf):
    ax.text(p["kg"] + 25, i, f"{p['kg']:,.0f} kg ({p['pct']:.0%})", va="center", fontsize=8, color=INK)
ax.set_xlabel("kg CO2e per year (auditor)")
style(ax, "x")
ax.set_xlim(0, max(p["kg"] for p in pf) * 1.32)
figs["personal"] = save(fig, "fig09_personal_footprint.png")

# Fig: EMO savings by category
em = [e for e in emos if e["hh_kwh"] and e["hh_kwh"] > 0]
em = sorted(em, key=lambda e: e["hh_kwh"])
catc = {"I": SER[0], "II": SER[1], "III": SER[2]}
fig, ax = plt.subplots(figsize=(7.2, 4.6))
ax.barh([e["ref"] for e in em], [e["hh_kwh"] for e in em], 0.6, color=[catc[e["cat"]] for e in em])
for i, e in enumerate(em):
    ax.text(e["hh_kwh"] + 5, i, f"{e['hh_kwh']:,.0f}", va="center", fontsize=7.8, color=INK)
from matplotlib.patches import Patch  # noqa: E402
ax.legend(handles=[Patch(color=catc[c], label=f"Category {c}") for c in ("I", "II", "III")], loc="lower right", fontsize=8)
ax.set_xlabel("Household share of saving, kWh per year")
style(ax, "x")
ax.set_xlim(0, max(e["hh_kwh"] for e in em) * 1.15)
figs["emo"] = save(fig, "fig10_emo_savings.png")

# Fig: sensitivity tornado
sens = [s for s in D["sensitivity"] if s["i"] != 0]
sens = sorted(sens, key=lambda s: abs(s["chg"]))
fig, ax = plt.subplots(figsize=(7.2, 4.2))
ax.barh([s["name"] for s in sens], [s["chg"] for s in sens], 0.55, color=[POS if s["chg"] > 0 else NEG for s in sens])
for i, s in enumerate(sens):
    off = 0.02 if s["chg"] >= 0 else -0.02
    ax.text(s["chg"] + off, i, f"{s['chg']:+.0%}", va="center", ha="left" if s["chg"] >= 0 else "right", fontsize=7.8, color=INK)
ax.axvline(0, color=INK2, linewidth=0.8)
ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
ax.set_xlabel("Change in household Scope 2 emissions vs base case")
style(ax, "x")
lo = min(s["chg"] for s in sens)
hi = max(s["chg"] for s in sens)
ax.set_xlim(lo - 0.12, hi + 0.15)
ax.legend(handles=[Patch(color=NEG, label="Lower than base"), Patch(color=POS, label="Higher than base")], loc="lower right", fontsize=8)
figs["sens"] = save(fig, "fig11_sensitivity.png")

# Fig: benchmark ratios
bms = [("Room EUI (kWh/m²)", k["eui_ratio"]), ("Electricity per person", k["pp_ratio"]),
       ("Fresh water per person", k["water_lpd"] / k["bm_water"]), ("Floor area per person", k["floor_pp"] / k["bm_floor_pp"]),
       ("Lighting power density", k["lpd"] / k["bec_lpd"])]
fig, ax = plt.subplots(figsize=(7.2, 2.5))
ax.barh([b[0] for b in bms][::-1], [b[1] for b in bms][::-1], 0.5, color=SER[0])
for i, b in enumerate(bms[::-1]):
    ax.text(b[1] + 0.02, i, f"{b[1]:.2f}×", va="center", fontsize=8, color=INK)
ax.axvline(1.0, color=INK2, linewidth=0.9, linestyle="-")
ax.text(1.0, len(bms) - 0.35, " benchmark = 1.0", fontsize=7.6, color=INK2, va="bottom")
ax.set_xlabel("Ratio to Hong Kong benchmark")
style(ax, "x")
ax.set_xlim(0, max(b[1] for b in bms) * 1.25)
figs["bench"] = save(fig, "fig12_benchmarks.png")

D["figs"] = {key: dict(file=name, size=Image.open(os.path.join(FIG, name)).size) for key, name in figs.items()}

# ---------------- photos
PHOTO_DIR = os.path.join(os.path.dirname(os.path.abspath(FIG)), "..", "photos")
PHOTO_DIR = os.path.normpath(PHOTO_DIR)
D["photo_dir"] = PHOTO_DIR
D["photos"] = {}
for f in sorted(os.listdir(PHOTO_DIR)):
    if f.lower().endswith(".jpg"):
        D["photos"][f] = Image.open(os.path.join(PHOTO_DIR, f)).size

# the AC specification screenshot is a tall list; the report shows only its energy-label half so it stays legible
crop = Image.open(os.path.join(PHOTO_DIR, "03a_ac_specification.jpg"))
crop = crop.crop((0, int(crop.height * 0.425), crop.width, crop.height))
crop_path = os.path.join(os.path.dirname(os.path.abspath(OUT)), "photo03a_energy_label_crop.jpg")
crop.save(crop_path, quality=92)
D["crops"] = {"03a": {"path": crop_path, "size": crop.size}}

# workbook statistics quoted in the report
wbf = load_workbook(WB)
D["wb_stats"] = {
    "sheets": len(wbf.sheetnames),
    "formulas": sum(1 for ws in wbf for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")),
}

with open(OUT, "w") as fh:
    json.dump(D, fh, indent=1, default=str)
print("figures:", len(figs), "->", FIG, "| data ->", OUT)

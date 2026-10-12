"""Build the Home Energy & Carbon Audit calculation workbook (live Excel formulas).

Usage:  python3 build_workbook.py <output.xlsx>
Then recalculate with LibreOffice (recalc.py) so cached values exist for viewers.
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

from audit_inputs import (AUDIT, DAYS_IN_MONTH, FLIGHTS, HKO_MEAN_T, HOT_MONTHS, INPUTS,
                          MONTHS, SOURCES)

OUT = sys.argv[1] if len(sys.argv) > 1 else "Home_Energy_Carbon_Audit_Workbook.xlsx"

ARIAL = "Arial"
F_IN = Font(name=ARIAL, size=10, color="0000FF")
F_CALC = Font(name=ARIAL, size=10, color="000000")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="1F3864")
F_SUB = Font(name=ARIAL, size=9, italic=True, color="595959")
F_SEC = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
F_HDR = Font(name=ARIAL, size=10, bold=True, color="1F3864")
F_NAME = Font(name=ARIAL, size=8, color="7F7F7F")
FILL_SEC = PatternFill("solid", fgColor="1F4E5F")
FILL_HDR = PatternFill("solid", fgColor="DCE6F0")
FILL_ASSUME = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOPWRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

KWH = '#,##0.0'
KWH0 = '#,##0'
PCT = '0.0%'
HKD = '"HK$"#,##0'
HKD2 = '"HK$"#,##0.00'

wb = Workbook()
wb.properties.creator = AUDIT["auditor"]
wb.properties.title = "Home Energy Audit [2] - calculation workbook"
NAMES = {}


def add_name(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
    NAMES[name] = ref


def put(ws, ref, value, kind="calc", fmt=None, bold=False, fill=None, wrap=False, border=True, align=None):
    c = ws[ref]
    c.value = value
    font = {"in": F_IN, "assume": F_IN, "calc": F_CALC, "link": F_LINK, "text": F_CALC,
            "note": F_SUB, "bold": F_BOLD, "name": F_NAME}[kind]
    if bold:
        font = Font(name=ARIAL, size=10, bold=True, color=font.color)
    c.font = font
    if kind == "assume":
        c.fill = FILL_ASSUME
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = TOPWRAP
    if align is not None:
        c.alignment = align
    if border and kind != "note":
        c.border = BOX
    return c


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def section(ws, row, text, first_col, last_col):
    for col in range(first_col, last_col + 1):
        ws.cell(row, col).fill = FILL_SEC
    c = ws.cell(row, first_col, text)
    c.font = F_SEC


def header(ws, row, labels, first_col=1, height=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row, first_col + i, lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.border = BOX
        c.alignment = CENTER
    if height:
        ws.row_dimensions[row].height = height


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def fmt_for(v):
    if isinstance(v, int):
        return KWH0
    s = repr(v)
    dec = len(s.split(".")[1]) if "." in s else 0
    return "0." + "0" * min(max(dec, 1), 5)


# ======================================================================================
# README (filled at the end once names exist)
ws_readme = wb.active
ws_readme.title = "README"

# ======================================================================================
# INPUTS
ws = wb.create_sheet("Inputs")
title(ws, "Inputs & assumptions",
      "Blue = input value; yellow fill = assumed typical value (flagged in the report); every formula in the workbook points here.")
header(ws, 4, ["Name", "Parameter", "Value", "Unit", "Data quality", "Source", "Notes"], height=30)
row, cur = 5, None
for name, sec, label, val, unit, q, src, note in INPUTS:
    if sec != cur:
        section(ws, row, sec, 1, 7)
        row += 1
        cur = sec
    put(ws, f"A{row}", name, "name")
    put(ws, f"B{row}", label, "text", wrap=True)
    put(ws, f"C{row}", val, "assume" if q == "Assumed" else "in", fmt=fmt_for(val))
    put(ws, f"D{row}", unit, "text")
    put(ws, f"E{row}", q, "text")
    put(ws, f"F{row}", src, "text")
    put(ws, f"G{row}", note, "text", wrap=True)
    add_name(name, f"Inputs!$C${row}")
    row += 1
last_input_row = row - 1

# Climate table (HKO normals) at I4:L16
header(ws, 4, ["Month", "Days", "Hot season (1 = May-Oct)", "HKO mean air temp (°C)"], first_col=9)
for i in range(12):
    r = 5 + i
    put(ws, f"I{r}", MONTHS[i], "text")
    put(ws, f"J{r}", DAYS_IN_MONTH[i], "in", KWH0)
    put(ws, f"K{r}", 1 if (i + 1) in HOT_MONTHS else 0, "in", KWH0)
    put(ws, f"L{r}", HKO_MEAN_T[i], "in", "0.0")
put(ws, "I17", "Year", "bold")
put(ws, "J17", "=SUM(J5:J16)", "calc", KWH0, bold=True)
put(ws, "K17", "=SUM(K5:K16)", "calc", KWH0, bold=True)
put(ws, "L17", "=SUMPRODUCT(J5:J16,L5:L16)/J17", "calc", "0.0", bold=True)
put(ws, "I18", "S8: HKO 1991-2020 normals (HKO HQ), day-weighted from 10-day & half-monthly normals; annual mean 23.5 °C.", "note")
widths(ws, {"A": 15, "B": 52, "C": 11, "D": 13, "E": 16, "F": 8, "G": 46, "H": 3, "I": 9, "J": 7, "K": 13, "L": 13})
ws.freeze_panes = "C5"

# ======================================================================================
# AC BASELINE
ws = wb.create_sheet("AC_Baseline")
title(ws, "Room air-conditioner baseline (Midea MW-12CRF8E, Octopus pre-paid)",
      "Top-down = reported Octopus spend ÷ electricity price.  Bottom-up = HK Energy Label average input power (362 kWh per 1,200 h) × reported operating hours.")
header(ws, 5, ["Month", "Days", "Season", "Octopus AC spend (HK$)", "Top-down kWh (spend ÷ tariff)",
               "AC operating hours (reported)", "Bottom-up kWh (label kW × hours)",
               "Scenario: stated HK$120/month (kWh)"], height=45)
for i in range(12):
    r, ir = 6 + i, 5 + i
    put(ws, f"A{r}", MONTHS[i], "text")
    put(ws, f"B{r}", f"=Inputs!J{ir}", "link", KWH0)
    put(ws, f"C{r}", f'=IF(Inputs!K{ir}=1,"Hot","Cool")', "calc")
    put(ws, f"D{r}", f"=IF(Inputs!K{ir}=1,AC_SPEND_HOT,AC_SPEND_COOL)", "calc", HKD)
    put(ws, f"E{r}", f"=D{r}/TARIFF", "calc", KWH)
    put(ws, f"F{r}", f"=B{r}*IF(Inputs!K{ir}=1,AC_H_HOT,AC_H_COOL)", "calc", KWH0)
    put(ws, f"G{r}", f"=F{r}*AC_LABEL_KWH/LABEL_HOURS", "calc", KWH)
    put(ws, f"H{r}", "=AC_SPEND_STATED/TARIFF", "calc", KWH)
put(ws, "A18", "Year", "bold", fill=FILL_TOT)
for col, f in zip("BDEFGH", [KWH0, HKD, KWH, KWH0, KWH, KWH]):
    put(ws, f"{col}18", f"=SUM({col}6:{col}17)", "calc", f, bold=True, fill=FILL_TOT)
put(ws, "C18", "", "calc", fill=FILL_TOT)

header(ws, 5, ["Reconciliation & diagnostics", "Value", "Unit", "Comment"], first_col=10)
diag = [
    ("Label average input power (HK label basis)", "=AC_LABEL_KWH/LABEL_HOURS", "kW", "362 kWh / 1,200 h", "0.000"),
    ("Rated input power (nameplate)", "=AC_RATED_KW", "kW", "Full-load inverter input", "0.000"),
    ("Annual operating hours (reported pattern)", "=F18", "h/yr", "12 h/day hot, 4 h/day cool", KWH0),
    ("BASELINE: top-down annual AC energy", "=E18", "kWh/yr", "Used in the inventory", KWH),
    ("Bottom-up annual AC energy (label method)", "=G18", "kWh/yr", "Independent physical check", KWH),
    ("Difference, bottom-up vs top-down", "=K10/K9-1", "%", "Within ±10% = good agreement", PCT),
    ("Upper-bound scenario (stated HK$120 × 12)", "=H18", "kWh/yr", "Not reconciled with seasonal answers", KWH),
    ("Implied Octopus rate if bottom-up is right", "=D18/G18", "HK$/kWh", "Cf. CUHK 1.20 - HKU 2.31 (S25)", "0.00"),
    ("Average input power while running (top-down)", "=E18/F18", "kW", "Inverter part-load operation", "0.000"),
    ("Equivalent full-load hours at rated input", "=E18/AC_RATED_KW", "h/yr", "", KWH0),
    ("Hot-season (May-Oct) AC energy", '=SUMIF(C6:C17,"Hot",E6:E17)', "kWh/yr", "", KWH),
    ("Hot-season share of AC energy", "=K16/K9", "%", "", PCT),
    ("Cool-season (Nov-Apr) AC energy", "=K9-K16", "kWh/yr", "", KWH),
    ("Annual AC spend (both occupants)", "=D18", "HK$/yr", "HK$570 per person", HKD),
    ("Approx. cooling delivered (top-down × label CSPF)", "=K9*AC_CSPF", "kWh(th)/yr", "Seasonal efficiency basis", KWH0),
]
for j, (lab, f, unit, com, nf) in enumerate(diag):
    r = 6 + j
    put(ws, f"J{r}", lab, "text", wrap=True)
    put(ws, f"K{r}", f, "calc", nf, fill=FILL_KEY if r == 9 else None, bold=(r == 9))
    put(ws, f"L{r}", unit, "text")
    put(ws, f"M{r}", com, "text", wrap=True)
add_name("AC_KWH", "AC_Baseline!$K$9")
add_name("AC_BU", "AC_Baseline!$K$10")
add_name("AC_STATED", "AC_Baseline!$K$12")
add_name("AC_KWH_HOT", "AC_Baseline!$K$16")
add_name("AC_SPEND_YR", "AC_Baseline!$K$19")
add_name("AC_HOURS", "AC_Baseline!$K$8")
widths(ws, {"A": 8, "B": 7, "C": 8, "D": 13, "E": 14, "F": 13, "G": 15, "H": 15, "I": 3, "J": 40, "K": 11, "L": 10, "M": 30})
ws.freeze_panes = "A6"

ch = BarChart()
ch.type = "col"
ch.title = "Room AC energy by month: top-down vs bottom-up"
ch.y_axis.title = "kWh/month"
ch.add_data(Reference(ws, min_col=5, min_row=5, max_row=17), titles_from_data=True)
ch.add_data(Reference(ws, min_col=7, min_row=5, max_row=17), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=17))
ch.height, ch.width = 8, 17
ws.add_chart(ch, "A21")

# ======================================================================================
# HOT WATER
ws = wb.create_sheet("HotWater")
title(ws, "Shower hot water - Stiebel Eltron DHB-E 27 instantaneous heater (shared washroom)",
      "Energy = litres × 4.186 kJ/kg.K × (40 °C - cold-water inlet) ÷ 3,600 ÷ efficiency.  Inlet temperature taken as the HKO monthly mean air temperature (assumption).")
header(ws, 5, ["Month", "Days", "Cold-water inlet ≈ HKO mean (°C)", "Temperature rise to 40 °C (K)",
               "Shower water per person (L/day)", "Heater energy per person (kWh/day)",
               "Household heater energy (kWh/month)", "Heater power needed at full flow (kW)"], height=45)
for i in range(12):
    r, ir = 6 + i, 5 + i
    put(ws, f"A{r}", MONTHS[i], "text")
    put(ws, f"B{r}", f"=Inputs!J{ir}", "link", KWH0)
    put(ws, f"C{r}", f"=Inputs!L{ir}", "link", "0.0")
    put(ws, f"D{r}", f"=SHOWER_T-C{r}", "calc", "0.0")
    put(ws, f"E{r}", "=SHOWERS*SHOWER_MIN*SHOWER_FLOW", "calc", KWH0)
    put(ws, f"F{r}", f"=E{r}*CP_WATER*D{r}/3600/HEATER_EFF", "calc", "0.000")
    put(ws, f"G{r}", f"=F{r}*B{r}*OCC", "calc", KWH)
    put(ws, f"H{r}", f"=SHOWER_FLOW/60*CP_WATER*D{r}/HEATER_EFF", "calc", "0.0")
put(ws, "A18", "Year", "bold", fill=FILL_TOT)
put(ws, "B18", "=SUM(B6:B17)", "calc", KWH0, bold=True, fill=FILL_TOT)
put(ws, "C18", "=SUMPRODUCT(B6:B17,C6:C17)/B18", "calc", "0.0", bold=True, fill=FILL_TOT)
put(ws, "D18", "=SUMPRODUCT(B6:B17,D6:D17)/B18", "calc", "0.00", bold=True, fill=FILL_TOT)
put(ws, "E18", "=E6", "calc", KWH0, bold=True, fill=FILL_TOT)
put(ws, "F18", "=G18/B18/OCC", "calc", "0.000", bold=True, fill=FILL_TOT)
put(ws, "G18", "=SUM(G6:G17)", "calc", KWH, bold=True, fill=FILL_TOT)
put(ws, "H18", "=MAX(H6:H17)", "calc", "0.0", bold=True, fill=FILL_TOT)
header(ws, 5, ["Diagnostics", "Value", "Unit"], first_col=10)
hw_diag = [
    ("Heater rating at 380 V (DHB-E 27)", "=HEATER_KW", "kW", "0.0"),
    ("Peak power needed for one shower (January)", "=H18", "kW", "0.0"),
    ("Capacity headroom (rating ÷ need)", "=K6/K7", "×", "0.00"),
    ("Household shower energy", "=G18", "kWh/yr", KWH),
    ("Per person", "=G18/OCC", "kWh/yr", KWH),
    ("Shower water per person", "=E6*365/1000", "m3/yr", "0.0"),
    ("January vs July daily energy", "=F6/F12", "×", "0.00"),
]
for j, (lab, f, unit, nf) in enumerate(hw_diag):
    r = 6 + j
    put(ws, f"J{r}", lab, "text", wrap=True)
    put(ws, f"K{r}", f, "calc", nf)
    put(ws, f"L{r}", unit, "text")
add_name("HW_KWH", "HotWater!$G$18")
add_name("HW_AVG_DT", "HotWater!$D$18")
add_name("SHOWER_LPD", "HotWater!$E$6")
widths(ws, {"A": 8, "B": 7, "C": 14, "D": 13, "E": 13, "F": 14, "G": 14, "H": 14, "I": 3, "J": 38, "K": 10, "L": 8})
ws.freeze_panes = "A6"

# ======================================================================================
# INVENTORY
ws = wb.create_sheet("Inventory")
title(ws, "Equipment inventory and annual energy (EAC 2018 Step 1-2)",
      "Household = Room 1003B (2 occupants). Room loads count in full; shared washroom fixed loads at 2/4; shared kitchen fixed loads at 2/32; usage-based items (showers, cooking, laundry) at the household's own use.")
header(ws, 6, ["ID", "Space", "Item", "Make / model / evidence", "Power basis (W)", "Qty",
               "Use (h/day) or basis", "Item annual kWh", "Household share", "Household kWh/yr",
               "Per person kWh/yr", "End-use", "EAC 8.1(g)(iv) category", "Data quality", "Photo / note"], height=45)
inv = [
    ("R1", "Room", "Window air-conditioner", "Midea MW-12CRF8E, 1.5 HP inverter, R32, Grade 1", "=AC_RATED_KW*1000", 1, "Octopus spend ÷ tariff", "=AC_KWH", 1, "Air-conditioning", "Air-conditioning", "Reported spend + label", "03a, 03b"),
    ("R2", "Room", "Ceiling LED linear battens", "2 surface battens (no label visible)", "=CEIL_W", "=CEIL_N", "=CEIL_H", None, 1, "Lighting", "Lighting", "Assumed W, reported h", "05a"),
    ("R3", "Room", "Under-shelf LED desk lights", "Integrated in furniture", "=DESK_W", "=DESK_N", "=DESK_H", None, 1, "Lighting", "Lighting", "Assumed", "05b"),
    ("R4", "Room", "Oscillating desk fan", "ZW02-180, 20 W nameplate", "=FAN_W", 1, "=FAN_H", None, 1, "Air movement (fan)", "Air-conditioning", "Nameplate + reported", "06b"),
    ("R5", "Room", "Laptops (auditor + roommate)", "HP 65 W adapter TPN-CA17 (avg draw 50%)", "=LAPTOP_W*LAPTOP_LF", "=LAPTOP_N", "=LAPTOP_H", None, 1, "Plug loads", "Others", "Nameplate + assumed LF", "06a"),
    ("R6", "Room", "Phone chargers", "Typical 20 W fast charger", "=PHONE_W", "=PHONE_N", "=PHONE_H", None, 1, "Plug loads", "Others", "Assumed W, reported h", "-"),
    ("W1", "Washroom", "Instantaneous water heater (household showers)", "Stiebel Eltron DHB-E 27, 24.4 kW @380 V", "=HEATER_KW*1000", 1, "See HotWater sheet", "=HW_KWH", 1, "Hot water", "Others", "Physics model", "08a, 08e"),
    ("W2", "Washroom", "Window exhaust fan (runs 24 h)", "J.E.E. window fan (label not visible)", "=WR_FAN_W", 1, "=WR_FAN_H", None, "=OCC/WR_USERS", "Ventilation", "Air-conditioning", "Assumed W, reported h", "08c, 08d"),
    ("W3", "Washroom", "LED downlights (occupancy sensor)", "3 surface downlights", "=WR_LIGHT_W", "=WR_LIGHT_N", "=WR_LIGHT_H", None, "=OCC/WR_USERS", "Lighting", "Lighting", "Assumed W, reported h", "08b"),
    ("K1", "Kitchen", "Refrigerators", "Panasonic NR-BV361B, Grade 1, 263 kWh/yr label", None, "=FRIDGE_N", "Energy label", "=FRIDGE_N*FRIDGE_KWH", "=OCC/KIT_USERS", "Refrigeration", "Others", "Energy label", "09i"),
    ("K2", "Kitchen", "3-temperature water dispenser", "Aquatek AQ-1323T, 750 W heater + 160 W compressor", 910, 1, "1.5 kWh/day (assumed)", "=DISP_KWH_DAY*365", "=OCC/KIT_USERS", "Drinking water", "Others", "Assumed", "09a, 09b"),
    ("K3", "Kitchen", "Split air-conditioner", "Toshiba RAS-24J2KCV-HK, 7.0 kW, R410A, Grade 1 (688 kWh/1,200 h)", "=KAC_LABEL_KWH/LABEL_HOURS*1000", 1, "8 h/day hot, 4 h/day cool", "=KAC_LABEL_KWH/LABEL_HOURS*(KAC_H_HOT*DAYS_HOT+KAC_H_COOL*DAYS_COOL)", "=OCC/KIT_USERS", "Air-conditioning", "Air-conditioning", "Label + observed", "09f, 09g"),
    ("K4", "Kitchen", "Kitchen lighting", "Not photographed", "=KLIGHT_W", "=KLIGHT_N", "=KLIGHT_H", None, "=OCC/KIT_USERS", "Lighting", "Lighting", "Assumed", "-"),
    ("K5", "Kitchen", "Induction cooking (household)", "Philips HD4911/80, 2,100 W max, Grade 2 (88.1%)", "=COOK_KW*1000", "=OCC", "=COOK_SESS*COOK_H/7", "=COOK_SESS*52*COOK_H*COOK_KW*OCC", 1, "Cooking", "Others", "Reported use + assumed kW", "09c-09e"),
    ("L1", "Laundry", "Washing machine (household)", "Hall machine, cold wash 45 min", None, "=OCC", "1 load/week each", "=LAUNDRY_WK*52*WASH_KWH*OCC", 1, "Laundry", "Others", "Reported use + assumed kWh", "-"),
    ("L2", "Laundry", "Tumble dryer (household)", "Hall machine, 12 min per load", "=DRY_KW*1000", "=OCC", "12 min/load", "=LAUNDRY_WK*52*DRY_KW*DRY_MIN/60*OCC", 1, "Laundry", "Others", "Reported use + assumed kW", "-"),
]
for k, (iid, space, item, model, pw, qty, use, h, share, enduse, eac, q, photo) in enumerate(inv):
    r = 7 + k
    put(ws, f"A{r}", iid, "bold")
    put(ws, f"B{r}", space, "text")
    put(ws, f"C{r}", item, "text", wrap=True)
    put(ws, f"D{r}", model, "text", wrap=True)
    put(ws, f"E{r}", pw, "link" if isinstance(pw, str) else "in", KWH0)
    put(ws, f"F{r}", qty, "link" if isinstance(qty, str) else "in", "0")
    if isinstance(use, str) and use.startswith("="):
        put(ws, f"G{r}", use, "link", "0.0")
    else:
        put(ws, f"G{r}", use, "text", wrap=True)
    if h is None:
        h = f"=E{r}*F{r}*G{r}*365/1000"
    put(ws, f"H{r}", h, "link" if h in ("=AC_KWH", "=HW_KWH") else "calc", KWH)
    put(ws, f"I{r}", share, "calc" if isinstance(share, str) else "in", "0.0000")
    put(ws, f"J{r}", f"=H{r}*I{r}", "calc", KWH)
    put(ws, f"K{r}", f"=J{r}/OCC", "calc", KWH)
    put(ws, f"L{r}", enduse, "text")
    put(ws, f"M{r}", eac, "text")
    put(ws, f"N{r}", q, "text", wrap=True)
    put(ws, f"O{r}", photo, "text")
put(ws, "A23", "Total", "bold", fill=FILL_TOT)
for col in "BCDEFGHI":
    put(ws, f"{col}23", "", "calc", fill=FILL_TOT)
put(ws, "J23", "=SUM(J7:J22)", "calc", KWH, bold=True, fill=FILL_KEY)
put(ws, "K23", "=SUM(K7:K22)", "calc", KWH, bold=True, fill=FILL_KEY)

header(ws, 26, ["Space", "Household kWh/yr", "% of household", "Per person kWh/yr"], first_col=3)
for k, (lab, key) in enumerate([("Room 1003B (15 m\u00b2)", "Room"), ("Shared washroom", "Washroom"),
                                 ("Shared kitchen", "Kitchen"), ("Hall laundry", "Laundry")]):
    r = 27 + k
    put(ws, f"C{r}", lab, "text")
    put(ws, f"D{r}", f'=SUMIF($B$7:$B$22,"{key}",$J$7:$J$22)', "calc", KWH)
    put(ws, f"E{r}", f"=D{r}/$D$31", "calc", PCT)
    put(ws, f"F{r}", f"=D{r}/OCC", "calc", KWH)
put(ws, "C31", "Total household", "bold", fill=FILL_TOT)
put(ws, "D31", "=SUM(D27:D30)", "calc", KWH, bold=True, fill=FILL_TOT)
put(ws, "E31", "=SUM(E27:E30)", "calc", PCT, bold=True, fill=FILL_TOT)
put(ws, "F31", "=D31/OCC", "calc", KWH, bold=True, fill=FILL_TOT)
put(ws, "C32", "Outside the room (shared / usage-based)", "text", wrap=True)
put(ws, "D32", "=D31-D27", "calc", KWH)
put(ws, "E32", "=D32/D31", "calc", PCT)
put(ws, "C33", "Check: space total - inventory total (must be 0)", "note")
put(ws, "D33", "=ROUND(D31-J23,6)", "calc", "0.000")
add_name("HH_KWH", "Inventory!$J$23")
add_name("ROOM_KWH", "Inventory!$D$27")
add_name("SHARED_KWH", "Inventory!$D$32")
widths(ws, {"A": 6, "B": 10, "C": 26, "D": 34, "E": 10, "F": 6, "G": 13, "H": 11, "I": 10, "J": 11,
            "K": 11, "L": 15, "M": 15, "N": 18, "O": 9})
ws.freeze_panes = "D7"

# ======================================================================================
# MONTHLY
ws = wb.create_sheet("Monthly")
title(ws, "Monthly energy, EUI and emissions profile (typical baseline year)",
      "Feeds EAC 2018 §8.1(g)(iii) monthly EUI and EE-EAes Part 2 item 5.  Room-only = loads inside the 15 m² room.")
header(ws, 5, ["Month", "Days", "Season", "Room AC", "Desk fan", "Room lighting", "Room plug loads", "Hot water",
               "Washroom fan + lights (share)", "Kitchen AC (share)", "Kitchen fridges, dispenser, lights (share)",
               "Cooking", "Laundry", "Household total kWh", "Household total MJ", "Room-only kWh",
               "Room EUI (MJ/m²)", "Per person kWh", "kg CO2e (CLP 0.34)", "kg CO2e (default 0.7)", "Octopus AC spend (HK$)"], height=58)
for i in range(12):
    r, ir = 6 + i, 5 + i
    put(ws, f"A{r}", MONTHS[i], "text")
    put(ws, f"B{r}", f"=Inputs!J{ir}", "link", KWH0)
    put(ws, f"C{r}", f'=IF(Inputs!K{ir}=1,"Hot","Cool")', "calc")
    put(ws, f"D{r}", f"=AC_Baseline!E{6 + i}", "link", KWH)
    put(ws, f"E{r}", f"=Inventory!$J$10*B{r}/365", "link", KWH)
    put(ws, f"F{r}", f"=(Inventory!$J$8+Inventory!$J$9)*B{r}/365", "link", KWH)
    put(ws, f"G{r}", f"=(Inventory!$J$11+Inventory!$J$12)*B{r}/365", "link", KWH)
    put(ws, f"H{r}", f"=HotWater!G{6 + i}", "link", KWH)
    put(ws, f"I{r}", f"=(Inventory!$J$14+Inventory!$J$15)*B{r}/365", "link", KWH)
    put(ws, f"J{r}", f"=KAC_LABEL_KWH/LABEL_HOURS*B{r}*IF(Inputs!K{ir}=1,KAC_H_HOT,KAC_H_COOL)*Inventory!$I$18", "calc", KWH)
    put(ws, f"K{r}", f"=(Inventory!$J$16+Inventory!$J$17+Inventory!$J$19)*B{r}/365", "link", KWH)
    put(ws, f"L{r}", f"=Inventory!$J$20*B{r}/365", "link", KWH)
    put(ws, f"M{r}", f"=(Inventory!$J$21+Inventory!$J$22)*B{r}/365", "link", KWH)
    put(ws, f"N{r}", f"=SUM(D{r}:M{r})", "calc", KWH, bold=True)
    put(ws, f"O{r}", f"=N{r}*MJ_PER_KWH", "calc", KWH0)
    put(ws, f"P{r}", f"=SUM(D{r}:G{r})", "calc", KWH)
    put(ws, f"Q{r}", f"=P{r}*MJ_PER_KWH/ROOM_AREA", "calc", "0.00")
    put(ws, f"R{r}", f"=N{r}/OCC", "calc", KWH)
    put(ws, f"S{r}", f"=N{r}*EF_CLP", "calc", KWH)
    put(ws, f"T{r}", f"=N{r}*EF_DEFAULT", "calc", KWH)
    put(ws, f"U{r}", f"=AC_Baseline!D{6 + i}", "link", HKD)
put(ws, "A18", "Year", "bold", fill=FILL_TOT)
put(ws, "C18", "", "calc", fill=FILL_TOT)
for col in "BDEFGHIJKLMNOPRSTU":
    put(ws, f"{col}18", f"=SUM({col}6:{col}17)", "calc", HKD if col == "U" else KWH0 if col in "BO" else KWH, bold=True, fill=FILL_TOT)
put(ws, "Q18", "=P18*MJ_PER_KWH/ROOM_AREA", "calc", "0.0", bold=True, fill=FILL_KEY)
put(ws, "A20", "Check: monthly total - inventory total (must be 0)", "note")
put(ws, "N20", "=ROUND(N18-HH_KWH,6)", "calc", "0.000")
put(ws, "A21", "Check: room-only total - room inventory (must be 0)", "note")
put(ws, "P21", "=ROUND(P18-ROOM_KWH,6)", "calc", "0.000")
put(ws, "A22", "Since move-in: Jan-Sep 2026 (9 months, EAC 8.1(g)(i) - count back to first day of occupation)", "bold", border=False)
put(ws, "A23", "Jan-Sep", "bold", fill=FILL_KEY)
for col in "BDEFGHIJKLMNOPRSTU":
    put(ws, f"{col}23", f"=SUM({col}6:{col}14)", "calc", HKD if col == "U" else KWH0 if col in "BO" else KWH, bold=True, fill=FILL_KEY)
put(ws, "C23", "", "calc", fill=FILL_KEY)
put(ws, "Q23", "=P23*MJ_PER_KWH/ROOM_AREA", "calc", "0.0", bold=True, fill=FILL_KEY)
add_name("HIST_HH_KWH", "Monthly!$N$23")
add_name("HIST_ROOM_KWH", "Monthly!$P$23")
add_name("HIST_SPEND", "Monthly!$U$23")
add_name("ROOM_EUI_MJ", "Monthly!$Q$18")
widths(ws, {"A": 7, "B": 6, "C": 7, **{c: 10 for c in "DEFGHIJKLM"}, "N": 11, "O": 10, "P": 10, "Q": 10, "R": 9, "S": 10, "T": 10, "U": 11})
ws.freeze_panes = "D6"

ch = BarChart()
ch.type = "col"
ch.grouping = "stacked"
ch.overlap = 100
ch.title = "Household electricity by month and end-use (kWh)"
ch.y_axis.title = "kWh/month"
ch.add_data(Reference(ws, min_col=4, max_col=13, min_row=5, max_row=17), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=17))
ch.height, ch.width = 10, 24
ws.add_chart(ch, "A26")

# ======================================================================================
# BREAKDOWN
ws = wb.create_sheet("Breakdown")
title(ws, "Energy breakdown, EUI and benchmarks",
      "EAC 2018 §8.1(g)(ii)-(iv): EUI, monthly EUI and percentage breakdown by installation.")
section(ws, 4, "A. End-use breakdown - household (Room 1003B + allocated shared services)", 1, 6)
header(ws, 5, ["End-use", "kWh/yr", "% of total", "kWh/person/yr", "kg CO2e/yr (CLP)", "kg CO2e/yr (default 0.7)"], height=30)
enduses = ["Air-conditioning", "Hot water", "Cooking", "Plug loads", "Lighting", "Ventilation",
           "Air movement (fan)", "Laundry", "Drinking water", "Refrigeration"]
for k, eu in enumerate(enduses):
    r = 6 + k
    put(ws, f"A{r}", eu, "text")
    put(ws, f"B{r}", f'=SUMIF(Inventory!$L$7:$L$22,A{r},Inventory!$J$7:$J$22)', "link", KWH)
    put(ws, f"C{r}", f"=B{r}/$B$16", "calc", PCT)
    put(ws, f"D{r}", f"=B{r}/OCC", "calc", KWH)
    put(ws, f"E{r}", f"=B{r}*EF_CLP", "calc", KWH)
    put(ws, f"F{r}", f"=B{r}*EF_DEFAULT", "calc", KWH)
put(ws, "A16", "Total", "bold", fill=FILL_TOT)
for col, nf in zip("BCDEF", [KWH, PCT, KWH, KWH, KWH]):
    put(ws, f"{col}16", f"=SUM({col}6:{col}15)", "calc", nf, bold=True, fill=FILL_TOT)

section(ws, 19, "B. EAC 2018 §8.1(g)(iv) categories", 1, 6)
header(ws, 20, ["Category", "Room 1003B kWh/yr", "% (room)", "Household kWh/yr", "% (household)", "Room MJ/yr"], height=30)
cats = ["Air-conditioning", "Lighting", "Others", "Lift & escalator"]
for k, cat in enumerate(cats):
    r = 21 + k
    put(ws, f"A{r}", cat, "text")
    if cat == "Lift & escalator":
        put(ws, f"B{r}", 0, "in", KWH)
        put(ws, f"D{r}", 0, "in", KWH)
    else:
        put(ws, f"B{r}", f'=SUMIFS(Inventory!$J$7:$J$22,Inventory!$M$7:$M$22,A{r},Inventory!$B$7:$B$22,"Room")', "link", KWH)
        put(ws, f"D{r}", f'=SUMIF(Inventory!$M$7:$M$22,A{r},Inventory!$J$7:$J$22)', "link", KWH)
    put(ws, f"C{r}", f"=B{r}/$B$25", "calc", PCT)
    put(ws, f"E{r}", f"=D{r}/$D$25", "calc", PCT)
    put(ws, f"F{r}", f"=B{r}*MJ_PER_KWH", "calc", KWH0)
put(ws, "A25", "Total", "bold", fill=FILL_TOT)
for col, nf in zip("BCDEF", [KWH, PCT, KWH, PCT, KWH0]):
    put(ws, f"{col}25", f"=SUM({col}21:{col}24)", "calc", nf, bold=True, fill=FILL_TOT)
put(ws, "A26", "Lift & escalator: not applicable to the audited unit (hall lifts are outside the household boundary).", "note")

section(ws, 28, "C. Key performance indicators and Hong Kong benchmarks", 1, 6)
header(ws, 29, ["Indicator", "Value", "Unit", "Benchmark", "Benchmark basis", "Ratio to benchmark"], height=30)
kpis = [
    ("Room EUI (energy inside the 15 m² room)", "=ROOM_KWH/ROOM_AREA", "kWh/m²/yr", "=BM_HH_KWH*12/BM_HH_PERS/BM_FLOOR_PP", "Typical HK home: 275 kWh/month ÷ 3 persons ÷ 16 m²/person (S7, S23)", "0.0"),
    ("Room EUI", "=ROOM_KWH*MJ_PER_KWH/ROOM_AREA", "MJ/m²/yr", "=D30*MJ_PER_KWH", "As above, in MJ", "0.0"),
    ("Household electricity per person", "=HH_KWH/OCC", "kWh/person/yr", "=BM_HH_KWH*12/BM_HH_PERS", "Typical HK household (S7)", KWH0),
    ("Share of household energy used outside the room", "=SHARED_KWH/HH_KWH", "%", "", "Shared washroom, kitchen, laundry", PCT),
    ("Room lighting power density", "=(CEIL_N*CEIL_W+DESK_N*DESK_W)/ROOM_AREA", "W/m²", "=BEC_LPD", "BEC 2021 max LPD, dormitory (S24)", "0.0"),
    ("Floor area per occupant", "=ROOM_AREA/OCC", "m²/person", "=BM_FLOOR_PP", "Median HK per-capita floor area (S23)", "0.0"),
    ("Fresh water use per person", "=GHG_Inventory!B40", "L/person/day", "=BM_WATER", "HK domestic fresh water 2023/24 (S13)", KWH),
    ("AC share of room energy", "=Breakdown!B21/Breakdown!B25", "%", "", "Incl. desk fan", PCT),
    ("Hot water share of household energy", "=Breakdown!B7/Breakdown!B16", "%", "", "", PCT),
]
for k, (lab, f, unit, bm, basis, nf) in enumerate(kpis):
    r = 30 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", f, "calc", nf, fill=FILL_KEY)
    put(ws, f"C{r}", unit, "text")
    put(ws, f"D{r}", bm, "calc", nf)
    put(ws, f"E{r}", basis, "text", wrap=True)
    put(ws, f"F{r}", f'=IF(D{r}="","",B{r}/D{r})', "calc", "0.00")
widths(ws, {"A": 40, "B": 13, "C": 13, "D": 13, "E": 40, "F": 12})

pie = PieChart()
pie.title = "Household electricity by EAC category"
pie.add_data(Reference(ws, min_col=4, min_row=20, max_row=23), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=1, min_row=21, max_row=23))
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showVal = False
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showLegendKey = False
pie.height, pie.width = 8, 12
ws.add_chart(pie, "H4")
bar = BarChart()
bar.type = "bar"
bar.title = "Household electricity by end-use (kWh/yr)"
bar.add_data(Reference(ws, min_col=2, min_row=5, max_row=15), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=6, max_row=15))
bar.legend = None
bar.height, bar.width = 9, 14
ws.add_chart(bar, "H21")

# ======================================================================================
# GHG INVENTORY
ws = wb.create_sheet("GHG_Inventory")
title(ws, "Greenhouse gas inventory - HK GHG Guidelines for Buildings (EPD/EMSD 2010) format",
      "Reporting entity: household of Room 1003B (2 occupants). Physical boundary: the room plus allocated shares of the shared washroom and kitchen, and the household's hall-laundry use. Baseline year: typical 12 months (data Jan-Oct 2026).")
section(ws, 4, "SCOPE 1 - Direct emissions", 1, 8)
put(ws, "A5", "Table 1 - Stationary combustion", "bold", border=False)
header(ws, 6, ["Source", "Fuel", "Amount", "Unit", "CO2e factor", "kg CO2e/yr"])
put(ws, "A7", "Cooking and water heating", "text")
put(ws, "B7", "None - all-electric (induction hob, electric heater)", "text", wrap=True)
put(ws, "C7", 0, "in", KWH)
put(ws, "D7", "-", "text")
put(ws, "E7", 0, "in", "0.000")
put(ws, "F7", "=C7*E7", "calc", KWH)
put(ws, "A9", "Table 3 - HFC refrigerant (mass-balance method: inventory change)", "bold", border=False)
header(ws, 10, ["Equipment", "Refrigerant", "Cs (kg)", "Ci (kg)", "Cd (kg)", "Ce (kg)", "GWP", "kg CO2e/yr"])
for k, (eq, refr, gwp) in enumerate([("Room window AC", "R32 (HFC-32)", "=GWP_R32"), ("Kitchen split AC", "R-410A", "=GWP_R410A")]):
    r = 11 + k
    put(ws, f"A{r}", eq, "text")
    put(ws, f"B{r}", refr, "text")
    for col in "CDEF":
        put(ws, f"{col}{r}", 0, "in", "0.00")
    put(ws, f"G{r}", gwp, "link", KWH0)
    put(ws, f"H{r}", f"=(C{r}+D{r}-E{r}-F{r})*G{r}", "calc", KWH)
put(ws, "A13", "Occupants hold, buy and dispose of no refrigerant, so the Guideline mass-balance gives zero. A leak-rate screening estimate is added below so fugitive HFCs are not omitted.", "note")
put(ws, "A15", "Screening estimate - IPCC 2006 operating leak rate (included in Scope 1)", "bold", border=False)
header(ws, 16, ["Equipment", "Refrigerant", "Charge (kg)", "Leak rate /yr", "GWP", "Household share", "kg CO2e/yr"])
put(ws, "A17", "Room window AC", "text")
put(ws, "B17", "R32", "text")
put(ws, "C17", "=AC_CHARGE", "link", "0.00")
put(ws, "D17", "=LEAK", "link", PCT)
put(ws, "E17", "=GWP_R32", "link", KWH0)
put(ws, "F17", 1, "in", "0.0000")
put(ws, "G17", "=C17*D17*E17*F17", "calc", KWH)
put(ws, "A18", "Kitchen split AC", "text")
put(ws, "B18", "R-410A", "text")
put(ws, "C18", "=KAC_CHARGE", "link", "0.00")
put(ws, "D18", "=LEAK", "link", PCT)
put(ws, "E18", "=GWP_R410A", "link", KWH0)
put(ws, "F18", "=OCC/KIT_USERS", "calc", "0.0000")
put(ws, "G18", "=C18*D18*E18*F18", "calc", KWH)
put(ws, "A19", "Scope 1 total", "bold", fill=FILL_TOT)
put(ws, "G19", "=F7+H11+H12+G17+G18", "calc", KWH, bold=True, fill=FILL_KEY)

section(ws, 21, "SCOPE 2 - Energy indirect emissions (Table 5: purchased electricity, CLP Power)", 1, 8)
header(ws, 22, ["Facility / source", "Electricity (kWh/yr)", "CLP-specific factor (kg/kWh)", "Territory default (kg/kWh)",
                "kg CO2e/yr (CLP)", "kg CO2e/yr (default)"], height=45)
s2rows = [("Room 1003B (incl. Octopus-metered AC)", "=Inventory!D27"),
          ("Shared washroom (own showers + 2/4 of fan & lights)", "=Inventory!D28"),
          ("Shared kitchen (own cooking + 2/32 of fixed loads)", "=Inventory!D29"),
          ("Hall laundry (own use)", "=Inventory!D30")]
for k, (lab, f) in enumerate(s2rows):
    r = 23 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", f, "link", KWH)
    put(ws, f"C{r}", "=EF_CLP", "link", "0.000")
    put(ws, f"D{r}", "=EF_DEFAULT", "link", "0.000")
    put(ws, f"E{r}", f"=B{r}*C{r}", "calc", KWH)
    put(ws, f"F{r}", f"=B{r}*D{r}", "calc", KWH)
put(ws, "A27", "Scope 2 total", "bold", fill=FILL_TOT)
put(ws, "B27", "=SUM(B23:B26)", "calc", KWH, bold=True, fill=FILL_TOT)
put(ws, "E27", "=SUM(E23:E26)", "calc", KWH, bold=True, fill=FILL_KEY)
put(ws, "F27", "=SUM(F23:F26)", "calc", KWH, bold=True, fill=FILL_KEY)
put(ws, "A28", "Table 6 - Towngas: not applicable (no Towngas or LPG used).", "note")

section(ws, 30, "SCOPE 3 - Other indirect emissions (optional under the Guidelines)", 1, 8)
header(ws, 31, ["Quantity", "Value", "Unit", "Basis"])
w3 = [
    ("Shower water", "=SHOWER_LPD", "L/person/day", "HotWater sheet", "0.0"),
    ("Drinking water (dispenser)", "=DRINK_L", "L/person/day", "Reported", "0.0"),
    ("Basin, cooking and dish-washing", "=OTHER_L", "L/person/day", "Assumed", "0.0"),
    ("Laundry water", "=WASH_L*LAUNDRY_WK/7", "L/person/day", "60 L per weekly load", "0.0"),
    ("Paper purchased (Pi)", "=SHEETS_SEM*SEMS*SHEET_G/1000*OCC", "kg/yr (household)", "100 sheets/semester each", "0.00"),
    ("Paper recycled (Pr)", "=B36*RECYCLE", "kg/yr", "Half recycled (assumed)", "0.00"),
    ("Paper in storage at start / end (Ps, Pe)", 0, "kg", "Negligible", "0.00"),
    ("", None, "", "", None),
    ("Fresh water per person", "=SUM(B32:B35)", "L/person/day", "", "0.0"),
    ("Household fresh water", "=B40*365*OCC/1000", "m3/yr", "", "0.0"),
    ("Household seawater (toilet flushing)", "=FLUSH_L*365*OCC/1000", "m3/yr", "5 flushes x 6 L/person/day", "0.0"),
]
# rows 32..42  (B40 = fresh L/person/day, B41 = household m3, B42 = seawater m3)
for k, (lab, f, unit, basis, nf) in enumerate(w3):
    r = 32 + k
    if lab == "":
        continue
    put(ws, f"A{r}", lab, "text")
    put(ws, f"B{r}", f, "link" if isinstance(f, str) else "in", nf)
    put(ws, f"C{r}", unit, "text")
    put(ws, f"D{r}", basis, "text")
header(ws, 44, ["Guideline item", "Quantity", "Unit", "Emission factor", "Factor unit", "kg CO2e/yr", "Factor basis"], height=30)
s3 = [
    ("Table 7 - Paper waste to landfill", "=B38+B36-B37-B38", "kg", "=PAPER_EF", "kg CO2e/kg", "Guideline default 4.8"),
    ("Table 8 - Fresh water (WSD)", "=B41", "m3", "=WSD_FRESH*EF_DEFAULT", "kg CO2e/m3", "0.624 kWh/m3 × 0.7"),
    ("Table 9 - Sewage (DSD)", "=B41*SEWER_FRAC", "m3", "=DSD_KWH*EF_DEFAULT", "kg CO2e/m3", "0.29 kWh/m3 × 0.7"),
    ("Seawater for flushing (WSD, optional)", "=B42", "m3", "=WSD_SALT*EF_DEFAULT", "kg CO2e/m3", "0.362 kWh/m3 × 0.7"),
]
for k, (lab, q, unit, ef, efu, basis) in enumerate(s3):
    r = 45 + k
    put(ws, f"A{r}", lab, "text")
    put(ws, f"B{r}", q, "calc", "0.00")
    put(ws, f"C{r}", unit, "text")
    put(ws, f"D{r}", ef, "link", "0.000")
    put(ws, f"E{r}", efu, "text")
    put(ws, f"F{r}", f"=B{r}*D{r}", "calc", KWH)
    put(ws, f"G{r}", basis, "text")
put(ws, "A49", "Scope 3 total", "bold", fill=FILL_TOT)
put(ws, "F49", "=SUM(F45:F48)", "calc", KWH, bold=True, fill=FILL_KEY)

section(ws, 51, "SUMMARY - household (Room 1003B, 2 occupants)", 1, 8)
header(ws, 52, ["Scope", "kg CO2e/yr (CLP-specific)", "kg CO2e/yr (territory default)", "t CO2e/yr (CLP)",
                "% of total (CLP)", "kg CO2e/person/yr (CLP)", "kg CO2e/person/yr (default)"], height=45)
summ = [("Scope 1 - direct (refrigerant screening)", "=G19", "=G19"),
        ("Scope 2 - purchased electricity", "=E27", "=F27"),
        ("Scope 3 - water, sewage, paper", "=F49", "=F49")]
for k, (lab, a, b) in enumerate(summ):
    r = 53 + k
    put(ws, f"A{r}", lab, "text")
    put(ws, f"B{r}", a, "calc", KWH)
    put(ws, f"C{r}", b, "calc", KWH)
    put(ws, f"D{r}", f"=B{r}/1000", "calc", "0.000")
    put(ws, f"E{r}", f"=B{r}/$B$56", "calc", PCT)
    put(ws, f"F{r}", f"=B{r}/OCC", "calc", KWH)
    put(ws, f"G{r}", f"=C{r}/OCC", "calc", KWH)
put(ws, "A56", "TOTAL", "bold", fill=FILL_TOT)
for col, nf in zip("BCDEFG", [KWH, KWH, "0.000", PCT, KWH, KWH]):
    put(ws, f"{col}56", f"=SUM({col}53:{col}55)", "calc", nf, bold=True, fill=FILL_KEY)
put(ws, "A57", "Emission intensity of the room (Scope 2, CLP basis)", "text")
put(ws, "B57", "=E23/ROOM_AREA", "calc", KWH)
put(ws, "C57", "kg CO2e/m²/yr", "text")
add_name("S1_KG", "GHG_Inventory!$B$53")
add_name("S2_CLP", "GHG_Inventory!$B$54")
add_name("S2_DEF", "GHG_Inventory!$C$54")
add_name("S3_KG", "GHG_Inventory!$B$55")
add_name("TOT_CLP", "GHG_Inventory!$B$56")
add_name("TOT_DEF", "GHG_Inventory!$C$56")
widths(ws, {"A": 44, "B": 15, "C": 15, "D": 14, "E": 14, "F": 15, "G": 18, "H": 12})

bar = BarChart()
bar.type = "col"
bar.title = "Household emissions by scope (kg CO2e/yr)"
bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=52, max_row=55), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=53, max_row=55))
bar.height, bar.width = 8, 16
ws.add_chart(bar, "J4")

# ======================================================================================
# LIFESTYLE (Scope 3 extension, per auditor)
ws = wb.create_sheet("Lifestyle_Scope3")
title(ws, "Daily-activity extension - the auditor's personal Scope 3 footprint",
      "Reported separately from the building inventory (GHG Protocol / ISO 14064-1 Scope 3 categories). Per auditor only, not per household.")
header(ws, 5, ["Activity", "Quantity", "Unit", "Emission factor", "Factor unit", "kg CO2e/yr", "Basis"], height=30)
put(ws, "A6", "MTR commute (18 km x 2 x 6 days/week)", "text")
put(ws, "B6", "=COMMUTE_KM_WK*WEEKS", "calc", KWH0)
put(ws, "C6", "p-km", "text")
put(ws, "D6", "=MTR_EF", "link", "0.0000")
put(ws, "E6", "kg CO2e/p-km", "text")
put(ws, "F6", "=B6*D6", "calc", KWH)
put(ws, "G6", "MTR CDP 2022 (S17); 48 weeks", "text")
put(ws, "A7", "Other MTR trips (22 km x 2 x 3/month)", "text")
put(ws, "B7", "=HOME_KM_MONTH*MONTHS_IN_HK", "calc", KWH0)
put(ws, "C7", "p-km", "text")
put(ws, "D7", "=MTR_EF", "link", "0.0000")
put(ws, "E7", "kg CO2e/p-km", "text")
put(ws, "F7", "=B7*D7", "calc", KWH)
put(ws, "G7", "11 months in HK", "text")
put(ws, "A8", "Flights (4 legs, table below)", "text")
put(ws, "B8", "=D18", "calc", KWH0)
put(ws, "C8", "p-km", "text")
put(ws, "D8", "=F8/B8", "calc", "0.0000")
put(ws, "E8", "kg CO2e/p-km", "text")
put(ws, "F8", "=F18", "calc", KWH)
put(ws, "G8", "DESNZ 2025 economy incl. RF (S18)", "text")
put(ws, "A9", "Diet - mixed (medium meat-eater)", "text")
put(ws, "B9", 365, "in", KWH0)
put(ws, "C9", "days", "text")
put(ws, "D9", "=DIET_EF", "link", "0.00")
put(ws, "E9", "kg CO2e/day", "text")
put(ws, "F9", "=B9*D9", "calc", KWH)
put(ws, "G9", "Scarborough et al. 2014 (S19)", "text")
put(ws, "A10", "Lifestyle total", "bold", fill=FILL_TOT)
put(ws, "F10", "=SUM(F6:F9)", "calc", KWH, bold=True, fill=FILL_KEY)
header(ws, 13, ["Flight leg", "Great-circle km", "Haul", "Distance incl. 8% uplift (km)", "Factor (kg CO2e/p-km)", "kg CO2e"], height=30)
for k, (leg, km, band) in enumerate(FLIGHTS):
    r = 14 + k
    put(ws, f"A{r}", leg, "text")
    put(ws, f"B{r}", km, "in", KWH0)
    put(ws, f"C{r}", "Long-haul" if band == "LONG" else "Short-haul", "in")
    put(ws, f"D{r}", f"=B{r}*FL_UPLIFT", "calc", KWH0)
    put(ws, f"E{r}", f'=IF(C{r}="Long-haul",EF_LONG,EF_SHORT)', "calc", "0.00000")
    put(ws, f"F{r}", f"=D{r}*E{r}", "calc", KWH)
put(ws, "A18", "Flights total", "bold", fill=FILL_TOT)
put(ws, "B18", "=SUM(B14:B17)", "calc", KWH0, bold=True, fill=FILL_TOT)
put(ws, "D18", "=SUM(D14:D17)", "calc", KWH0, bold=True, fill=FILL_TOT)
put(ws, "F18", "=SUM(F14:F17)", "calc", KWH, bold=True, fill=FILL_TOT)
header(ws, 21, ["Personal footprint (auditor)", "kg CO2e/yr", "% of personal total"], height=30)
pf = [("Home energy & water (household inventory ÷ 2)", "=TOT_CLP/OCC"),
      ("MTR travel", "=F6+F7"), ("Flights", "=F8"), ("Diet", "=F9")]
for k, (lab, f) in enumerate(pf):
    r = 22 + k
    put(ws, f"A{r}", lab, "text")
    put(ws, f"B{r}", f, "calc", KWH)
    put(ws, f"C{r}", f"=B{r}/$B$26", "calc", PCT)
put(ws, "A26", "Personal total", "bold", fill=FILL_TOT)
put(ws, "B26", "=SUM(B22:B25)", "calc", KWH, bold=True, fill=FILL_KEY)
put(ws, "C26", "=SUM(C22:C25)", "calc", PCT, bold=True, fill=FILL_TOT)
put(ws, "A27", "HK per-capita emissions 2023 (territorial, S10)", "text")
put(ws, "B27", "=BM_HK_T*1000", "link", KWH)
put(ws, "A28", "Ratio: personal total / HK per-capita (different scope - see report)", "text")
put(ws, "B28", "=B26/B27", "calc", "0.00")
add_name("LIFE_KG", "Lifestyle_Scope3!$F$10")
add_name("PERSONAL_KG", "Lifestyle_Scope3!$B$26")
widths(ws, {"A": 46, "B": 13, "C": 13, "D": 16, "E": 14, "F": 12, "G": 32})
bar = BarChart()
bar.type = "bar"
bar.title = "Auditor's personal footprint (kg CO2e/yr)"
bar.add_data(Reference(ws, min_col=2, min_row=21, max_row=25), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=22, max_row=25))
bar.legend = None
bar.height, bar.width = 7, 15
ws.add_chart(bar, "I4")

# ======================================================================================
# EMO REGISTER
ws = wb.create_sheet("EMO_Register")
title(ws, "Energy Management Opportunities (EAC 2018 §6.2 categories; EE-EAes Part 3 types)",
      "Cat I = housekeeping, no cost; Cat II = low-cost operational changes; Cat III = higher capital cost.  Types: LG lighting, AC air-conditioning, EL electrical, LE lift/escalator, Others.  Capital costs are indicative and must be confirmed by quotation.  Life-cycle columns (P-W) follow EAC §7.5.4: service life and degradation of savings.")
header(ws, 5, ["Ref", "Cat.", "Type", "Responsible party", "Energy management opportunity", "Basis of saving estimate",
               "Facility saving (kWh/yr)", "Household share (kWh/yr)", "Household saving (HK$/yr)",
               "Indicative capital cost (HK$)", "Simple payback (yr)", "Service life (yr)",
               "Facility saving (kg CO2e/yr)", "Household saving (kg CO2e/yr)", "Recommendation",
               "Affected end-use before the EMO (kWh/yr, facility)", "Share of that end-use saved",
               "Facility saving (HK$/yr)", "NPV over service life (HK$)", "Savings-to-investment ratio",
               "Cost of saved energy (HK$/kWh)", "Life-cycle saving (t CO2e)", "Net cost per t CO2e avoided (HK$/t)"], height=70)
AC_DJF = "(AC_Baseline!E6+AC_Baseline!E7+AC_Baseline!E17)"
emos = [
    # ref, cat, type, responsible, description, basis, facility kWh, household share, cost, life, recommendation, affected end-use baseline (facility)
    ("EMO-01", "I", "AC", "Occupants", "Raise the AC set-point from 25 °C to 26 °C and keep the 20 W desk fan running for air movement", "3% of AC energy per +1 °C (S14)", "=AC_KWH*DEG_SAVING*(26-AC_SETPOINT)", None, 0, "-", "Implement now", "=AC_KWH"),
    ("EMO-02", "I", "AC", "Occupants", "Use the AC's built-in timer to switch off 2 h before waking in hot months (12 to 10 h/day)", "2/12 of hot-season AC energy", "=AC_KWH_HOT*2/AC_H_HOT", None, 0, "-", "Implement now", "=AC_KWH_HOT"),
    ("EMO-03", "I", "AC", "Occupants", "Clean the AC filter fortnightly in summer, wipe the soiled louvres, close the curtain against late-afternoon NW sun", "5% of AC energy (S26, S32)", "=AC_KWH*0.05", None, 0, "-", "Implement now", "=AC_KWH"),
    ("EMO-04", "I", "AC", "Occupants", "Cool-season 'fan first': no AC cooling in December-February (HKO mean 16.5-18.2 °C); use the desk fan and open window, Dry mode only on humid days", "Room-AC energy of Dec, Jan and Feb avoided", f"={AC_DJF}", None, 0, "-", "Implement now", f"={AC_DJF}"),
    ("EMO-05", "I", "LG", "Occupants", "Use the desk task lights instead of both ceiling battens for 3 of the 6 lit hours", "2 x 20 W x 3 h/day avoided", "=CEIL_N*CEIL_W*3*365/1000", None, 0, "-", "Implement now", "=Inventory!J8"),
    ("EMO-06", "I", "Others", "Occupants", "Shorten showers from 5 to 4 minutes (phone timer or a 4-minute song)", "20% of shower heating energy", "=HW_KWH*(1-4/SHOWER_MIN)", None, 0, "-", "Implement now", "=HW_KWH"),
    ("EMO-07", "I", "Others", "Occupants", "Lower the delivered shower temperature from 40 °C to 38 °C (turn the heater dial down rather than mixing cold water)", "2 K of the annual-average temperature rise", "=HW_KWH*2/HW_AVG_DT", None, 0, "-", "Implement now", "=HW_KWH"),
    ("EMO-08", "I", "EL", "Occupants", "Laptop power management: lower brightness, sleep after 10 min idle, battery-saver mode", "15% of laptop energy (assumed)", "=Inventory!J11*0.15", None, 0, "-", "Implement now", "=Inventory!J11"),
    ("EMO-09", "I", "Others", "Occupants", "Cook with lids, match pan to the induction zone, batch-cook two meals per session", "20% of cooking energy (assumed)", "=Inventory!J20*0.2", None, 0, "-", "Implement now", "=Inventory!J20"),
    ("EMO-10", "I", "Others", "Occupants", "Air-dry laundry on a rack by the open window instead of tumble-drying (keep the dryer for bulky items in humid spells)", "80% of tumble-dryer energy avoided", "=Inventory!J22*DRYER_DISPLACE", None, 0, "-", "Implement now", "=Inventory!J22"),
    ("EMO-11", "II", "AC", "Hall (Student Residence Office)", "Washroom exhaust fan: replace 24 h running with an occupancy / humidistat run-on control (~8 h/day)", "16 of 24 fan-hours avoided", "=WR_FAN_W*(WR_FAN_H-8)*365/1000", "=G{r}*OCC/WR_USERS", 800, 10, "Recommend to hall", "=WR_FAN_W*WR_FAN_H*365/1000"),
    ("EMO-12", "II", "AC", "Hall (Student Residence Office)", "Kitchen AC: 7-day timer + occupancy interlock, set-point locked at 25.5 °C (8 to 5 h/day hot, 4 to 2 h/day cool)", "Label 0.573 kW x hours avoided", "=KAC_LABEL_KWH/LABEL_HOURS*((KAC_H_HOT-5)*DAYS_HOT+(KAC_H_COOL-2)*DAYS_COOL)", "=G{r}*OCC/KIT_USERS", 1000, 10, "Recommend to hall", "=Inventory!H18"),
    ("EMO-13", "II", "Others", "Hall (Student Residence Office)", "Water dispenser: enable night set-back / timer on hot and ice tanks (00:00-06:00)", "25% of dispenser energy (assumed)", "=DISP_KWH_DAY*365*0.25", "=G{r}*OCC/KIT_USERS", 200, 10, "Recommend to hall", "=Inventory!H17"),
    ("EMO-14", "II", "Others", "Hall (Student Residence Office)", "Fit a WELS Grade-1 low-flow showerhead (9 to 7 L/min) in the shared washroom", "22% of shower energy and water, 4 users", "=HW_KWH/OCC*WR_USERS*(1-7/SHOWER_FLOW)", "=HW_KWH*(1-7/SHOWER_FLOW)", 250, 5, "Recommend to hall", "=HW_KWH/OCC*WR_USERS"),
    ("EMO-15", "II", "LG", "Hall (Student Residence Office)", "Kitchen lighting: two ceiling occupancy (PIR) sensors so lights follow use (16 to 8 h/day)", "6 x 15 W x 8 h/day avoided", "=KLIGHT_N*KLIGHT_W*(KLIGHT_H-KLIGHT_SENSOR_H)*365/1000", "=G{r}*OCC/KIT_USERS", 600, 10, "Recommend to hall", "=Inventory!H19"),
    ("EMO-16", "II", "AC", "Hall (Student Residence Office)", "Annual professional clean of the window-AC coil and louvres (heavy soiling seen in photo 03b)", "3% of AC energy (S26)", "=AC_KWH*0.03", None, 400, 1, "Recommend on hygiene grounds; bundle with routine maintenance (EAC 7.6.4)", "=AC_KWH"),
    ("EMO-17", "II", "EL", "Occupants", "Buy a plug-in power meter (about HK$120) to measure the laptop, fan and AC and verify savings (EAC §7.6.3)", "Enabling measure - no saving claimed", 0, None, 120, 5, "Implement now (within HK$200 budget)", 0),
    ("EMO-18", "III", "Others", "Hall (CityU)", "Replace 24.4 kW instantaneous shower heaters with central air-source heat-pump water heating (COP about 3.5) and storage", "Electricity x (1 - 0.99/3.5), 4 users", "=HW_KWH/OCC*WR_USERS*(1-HEATER_EFF/3.5)", "=HW_KWH*(1-HEATER_EFF/3.5)", 20000, 15, "Further study (also cuts peak demand)", "=HW_KWH/OCC*WR_USERS"),
    ("EMO-19", "III", "AC", "Hall (CityU)", "Replace the Grade-1 window AC (CSPF 4.66) with a best-in-class split unit (CSPF about 7.5)", "AC x (1 - 4.66/7.5)", "=AC_KWH*(1-AC_CSPF/7.5)", None, 9000, 12, "Not viable - payback exceeds service life", "=AC_KWH"),
    ("EMO-20", "III", "AC", "Hall (CityU)", "Solar-control film on the NW single glazing (1.55 m²)", "2% of AC energy (assumed)", "=AC_KWH*0.02", None, 800, 10, "Not viable at present", "=AC_KWH"),
    ("EMO-21", "III", "Others", "Hall (CityU)", "Expand rooftop solar PV under the CLP Feed-in Tariff (the village already has 110 panels)", "Renewable-energy viability statement (EAC §7.4.4(d))", 0, None, 0, 20, "Further study by CityU", 0),
]
first_emo = 6
for k, (ref, cat, typ, resp, desc, basis, g, h, cost, life, rec, base) in enumerate(emos):
    r = first_emo + k
    put(ws, f"A{r}", ref, "bold")
    put(ws, f"B{r}", cat, "text", align=CENTER)
    put(ws, f"C{r}", typ, "text", align=CENTER)
    put(ws, f"D{r}", resp, "text", wrap=True)
    put(ws, f"E{r}", desc, "text", wrap=True)
    put(ws, f"F{r}", basis, "text", wrap=True)
    put(ws, f"G{r}", g, "calc" if isinstance(g, str) else "in", KWH)
    put(ws, f"H{r}", (h or "=G{r}").format(r=r), "calc", KWH)
    put(ws, f"I{r}", f"=H{r}*TARIFF", "calc", HKD)
    put(ws, f"J{r}", cost, "assume" if cost else "in", HKD)
    put(ws, f"K{r}", f'=IF(J{r}=0,0,IFERROR(J{r}/(G{r}*TARIFF),"n/a"))', "calc", "0.0")
    put(ws, f"L{r}", life, "in")
    put(ws, f"M{r}", f"=G{r}*EF_CLP", "calc", KWH)
    put(ws, f"N{r}", f"=H{r}*EF_CLP", "calc", KWH)
    put(ws, f"O{r}", rec, "text", wrap=True)
    put(ws, f"P{r}", base, "calc" if isinstance(base, str) else "in", KWH)
    put(ws, f"Q{r}", f'=IF(P{r}=0,"-",G{r}/P{r})', "calc", PCT)
    put(ws, f"R{r}", f"=G{r}*TARIFF", "calc", HKD)
    # NPV with savings declining by DEGRADE per year: PV = S x (1 - q^L)/(r + d), q = (1-d)/(1+r)
    put(ws, f"S{r}", f'=IF(J{r}=0,"-",R{r}*(1-((1-DEGRADE)/(1+DISC_RATE))^L{r})/(DISC_RATE+DEGRADE)-J{r})', "calc", HKD)
    put(ws, f"T{r}", f'=IF(J{r}=0,"-",(S{r}+J{r})/J{r})', "calc", "0.00")
    put(ws, f"U{r}", f'=IF(OR(J{r}=0,G{r}=0),"-",J{r}*DISC_RATE/(1-(1+DISC_RATE)^-L{r})/G{r})', "calc", HKD2)
    put(ws, f"V{r}", f'=IF(J{r}=0,"-",M{r}*(1-(1-DEGRADE)^L{r})/DEGRADE/1000)', "calc", "0.00")
    put(ws, f"W{r}", f'=IF(OR(J{r}=0,G{r}=0),"-",-S{r}/V{r})', "calc", HKD)
last_emo = first_emo + len(emos) - 1
rng = lambda col: f"${col}${first_emo}:${col}${last_emo}"
put(ws, f"A{last_emo + 1}", "Life-cycle columns: NPV = facility saving x (1 - ((1 - d)/(1 + r))^L)/(r + d) - cost, with r = DISC_RATE and d = DEGRADE; "
    "cost of saved energy = cost x r/(1 - (1 + r)^-L) ÷ annual kWh (compare with the HK$1.406/kWh tariff); "
    "net cost per t CO2e = -NPV ÷ life-cycle tonnes (negative = saves money).", "note")
s0 = last_emo + 3
header(ws, s0, ["Category", "No. of EMOs", "Household saving (kWh/yr)", "% of household electricity",
                "Household saving (kg CO2e/yr)", "Facility saving (kWh/yr)"], height=45)
for k, cat in enumerate(["I", "II", "III"]):
    r = s0 + 1 + k
    put(ws, f"A{r}", f"Category {cat}", "text")
    put(ws, f"B{r}", f'=COUNTIF({rng("B")},"{cat}")', "calc", "0")
    put(ws, f"C{r}", f'=SUMIF({rng("B")},"{cat}",{rng("H")})', "calc", KWH)
    put(ws, f"D{r}", f"=C{r}/HH_KWH", "calc", PCT)
    put(ws, f"E{r}", f'=SUMIF({rng("B")},"{cat}",{rng("N")})', "calc", KWH)
    put(ws, f"F{r}", f'=SUMIF({rng("B")},"{cat}",{rng("G")})', "calc", KWH)
tr = s0 + 4
put(ws, f"A{tr}", "All EMOs (indicative - savings interact, not fully additive)", "bold", fill=FILL_TOT)
for col, nf in zip("BCDEF", ["0", KWH, PCT, KWH, KWH]):
    put(ws, f"{col}{tr}", f"=SUM({col}{s0 + 1}:{col}{s0 + 3})", "calc", nf, bold=True, fill=FILL_TOT)
t0 = tr + 3
header(ws, t0, ["EE-EAes type", "No. of EMOs"])
for k, (typ, lab) in enumerate([("LG", "Lighting (LG)"), ("AC", "Air-conditioning (AC)"), ("EL", "Electrical (EL)"),
                                ("LE", "Lift/Escalator (LE)"), ("Others", "Others")]):
    r = t0 + 1 + k
    put(ws, f"A{r}", lab, "text")
    put(ws, f"B{r}", f'=COUNTIF({rng("C")},"{typ}")', "calc", "0")
add_name("EMO_CAT1_KWH", f"EMO_Register!$C${s0 + 1}")
add_name("EMO_CAT1_KG", f"EMO_Register!$E${s0 + 1}")
EMO_SUMMARY_ROW = s0
EMO_TYPE_ROW = t0
widths(ws, {"A": 9, "B": 5, "C": 7, "D": 16, "E": 46, "F": 26, "G": 11, "H": 11, "I": 11, "J": 11, "K": 9,
            "L": 8, "M": 11, "N": 11, "O": 22, "P": 13, "Q": 10, "R": 11, "S": 12, "T": 11, "U": 12, "V": 11, "W": 13})
ws.freeze_panes = "F6"
bar = BarChart()
bar.type = "bar"
bar.title = "Household saving by EMO (kWh/yr)"
bar.add_data(Reference(ws, min_col=8, min_row=5, max_row=last_emo), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=first_emo, max_row=last_emo))
bar.legend = None
bar.height, bar.width = 12, 15
ws.add_chart(bar, "Y5")

# ======================================================================================
# SENSITIVITY
ws = wb.create_sheet("Sensitivity")
title(ws, "Sensitivity analysis - which assumptions matter?",
      "Each scenario changes one input and recomputes the household total; Scope 2 uses the CLP factor unless the scenario changes the factor.")
header(ws, 5, ["#", "Scenario", "Value tested", "Room AC (kWh/yr)", "Household electricity (kWh/yr)",
               "Household Scope 2 (kg CO2e/yr)", "Change vs base"], height=45)
lap, cook, wash, dry = "Inventory!J11", "Inventory!J20", "Inventory!J21", "Inventory!J22"
kfix = "(Inventory!J16+Inventory!J17+Inventory!J18+Inventory!J19)"
scen = [
    ("Base case", "CLP 0.34; HK$1.406/kWh", "=AC_KWH", "=HH_KWH", "=E{r}*EF_CLP"),
    ("CLP 2024 emission factor", "=EF_CLP_PREV", "=AC_KWH", "=HH_KWH", "=E{r}*EF_CLP_PREV"),
    ("Territory-wide default factor", "=EF_DEFAULT", "=AC_KWH", "=HH_KWH", "=E{r}*EF_DEFAULT"),
    ("Octopus AC rate HK$1.20/kWh (CUHK level)", "=TARIFF_LOW", "=AC_SPEND_YR/TARIFF_LOW", "=HH_KWH-AC_KWH+D{r}", "=E{r}*EF_CLP"),
    ("Octopus AC rate HK$2.31/kWh (HKU level)", "=TARIFF_HIGH", "=AC_SPEND_YR/TARIFF_HIGH", "=HH_KWH-AC_KWH+D{r}", "=E{r}*EF_CLP"),
    ("Stated flat average HK$120/month", "=AC_SPEND_STATED", "=AC_STATED", "=HH_KWH-AC_KWH+D{r}", "=E{r}*EF_CLP"),
    ("Bottom-up AC (label 362 kWh/1,200 h)", "=AC_LABEL_KWH", "=AC_BU", "=HH_KWH-AC_KWH+D{r}", "=E{r}*EF_CLP"),
    ("Shower flow 7 L/min (Grade-1 head)", 7, "=AC_KWH", "=HH_KWH-HW_KWH*(1-C{r}/SHOWER_FLOW)", "=E{r}*EF_CLP"),
    ("Shower flow 12 L/min", 12, "=AC_KWH", "=HH_KWH-HW_KWH*(1-C{r}/SHOWER_FLOW)", "=E{r}*EF_CLP"),
    ("Laptop average draw 30% of adapter", 0.3, "=AC_KWH", f"=HH_KWH-{lap}*(1-C{{r}}/LAPTOP_LF)", "=E{r}*EF_CLP"),
    ("Laptop average draw 80% of adapter", 0.8, "=AC_KWH", f"=HH_KWH-{lap}*(1-C{{r}}/LAPTOP_LF)", "=E{r}*EF_CLP"),
    ("Roommate uses half the auditor's personal loads", 0.5, "=AC_KWH", f"=HH_KWH-(1-C{{r}})/2*({lap}+Inventory!J12+HW_KWH+{cook}+{wash}+{dry})", "=E{r}*EF_CLP"),
    ("Kitchen allocated over 20 active users (not 32)", 20, "=AC_KWH", f"=HH_KWH+{kfix}*(KIT_USERS/C{{r}}-1)", "=E{r}*EF_CLP"),
    ("One month away (11 of 12 months occupied)", "=11/12", "=AC_KWH*C{r}", f"=HH_KWH-(1-C{{r}})*(ROOM_KWH+HW_KWH+{cook}+{wash}+{dry})", "=E{r}*EF_CLP"),
]
for k, (lab, val, d, e, f) in enumerate(scen):
    r = 6 + k
    put(ws, f"A{r}", k, "text")
    put(ws, f"B{r}", lab, "text", wrap=True)
    put(ws, f"C{r}", val, "link" if isinstance(val, str) else "in", "0.000" if isinstance(val, float) or isinstance(val, str) else "0")
    put(ws, f"D{r}", d.format(r=r), "calc", KWH)
    put(ws, f"E{r}", e.format(r=r), "calc", KWH)
    put(ws, f"F{r}", f.format(r=r), "calc", KWH)
    put(ws, f"G{r}", f"=F{r}/$F$6-1", "calc", PCT)
put(ws, "B6", "Base case", "bold")
widths(ws, {"A": 4, "B": 44, "C": 12, "D": 13, "E": 15, "F": 15, "G": 11})
bar = BarChart()
bar.type = "bar"
bar.title = "Change in household Scope 2 vs base"
bar.add_data(Reference(ws, min_col=7, min_row=5, max_row=6 + len(scen) - 1), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=2, min_row=6, max_row=6 + len(scen) - 1))
bar.legend = None
bar.height, bar.width = 11, 16
ws.add_chart(bar, "I5")

# ======================================================================================
# DISTRICT SCALING (assignment part 3: Sha Tin District)
ws = wb.create_sheet("District_Scaling")
title(ws, "District scale-up - Sha Tin District (Ma On Shan lies in Sha Tin)",
      "Energy per capita and per square metre from this audit are scaled to the district with C&SD and Housing Department data; a composite factor links energy to floor area and number of residents.")
section(ws, 4, "A. Factors from the audit (household of Room 1003B)", 1, 4)
header(ws, 5, ["Factor", "Value", "Unit", "Basis"], height=30)
fa = [
    ("Electricity per capita (e_pc)", "=HH_KWH/OCC", "kWh/person/yr", "Household electricity ÷ 2 occupants", KWH),
    ("Electricity per m² of dwelling floor area (e_m2)", "=HH_KWH/ROOM_AREA", "kWh/m²/yr", "Household electricity ÷ 15 m² room (shared-facility energy included)", KWH),
    ("Room EUI (energy used inside the room only)", "=ROOM_KWH/ROOM_AREA", "kWh/m²/yr", "For reference (EAC EUI basis)", KWH),
    ("Floor area per person - audited household", "=ROOM_AREA/OCC", "m²/person", "", "0.0"),
    ("Area-driven energy: room AC + room lighting", "=Inventory!J7+Inventory!J8+Inventory!J9", "kWh/yr", "Loads that scale with conditioned and lit floor area", KWH),
    ("a = area coefficient", "=B10/ROOM_AREA", "kWh/m²/yr", "Area-driven energy ÷ floor area", "0.00"),
    ("Person-driven energy: all other loads", "=HH_KWH-B10", "kWh/yr", "Hot water, cooking, plug loads, fan, laundry, shared services", KWH),
    ("b = resident coefficient", "=B12/OCC", "kWh/person/yr", "Person-driven energy ÷ occupants", "0.0"),
    ("Area-driven share of household electricity", "=B10/HH_KWH", "%", "", PCT),
]
for k, (lab, f, unit, basis, nf) in enumerate(fa):
    r = 6 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", f, "calc", nf, fill=FILL_KEY if r in (6, 7, 11, 13) else None, bold=r in (6, 7, 11, 13))
    put(ws, f"C{r}", unit, "text")
    put(ws, f"D{r}", basis, "text", wrap=True)
add_name("DS_EPC", "District_Scaling!$B$6")
add_name("DS_EM2", "District_Scaling!$B$7")
add_name("DS_A", "District_Scaling!$B$11")
add_name("DS_B", "District_Scaling!$B$13")

section(ws, 16, "B. Sha Tin District population, households and domestic floor area", 1, 4)
header(ws, 17, ["Quantity", "Value", "Unit", "Basis / source"], height=30)
fb = [
    ("Population (land-based non-institutional), 2025", "=ST_POP", "persons", "C&SD district statistics 2025 (S33)", KWH0),
    ("Domestic households, 2025", "=ST_HH", "households", "S33", KWH0),
    ("Average household size", "=ST_AHS", "persons/household", "S33", "0.0"),
    ("Median floor area per person (Hong Kong, 2021 Census)", "=BM_FLOOR_PP", "m²/person", "S23", "0.0"),
    ("Method 1 - domestic floor area = population x median floor area per person", "=ST_POP*BM_FLOOR_PP", "m²", "", KWH0),
    ("   floor area per household (method 1)", "=B22/ST_HH", "m²/household", "", "0.0"),
    ("Housing Authority public rental households in Sha Tin", "=ST_PRH_HH", "households", "S35", KWH0),
    ("   public rental floor area = households x 14.1 m²/person x 2.6 persons", "=ST_PRH_HH*PRH_SPACE_PP*PRH_AHS", "m²", "Housing in Figures 2025 (S34)", KWH0),
    ("Other households (subsidised sale + private)", "=ST_HH-ST_PRH_HH", "households", "", KWH0),
    ("   subsidised-sale share of other households (2021 Census mix)", "=ST_SSF_2021/(ST_SSF_2021+ST_PRIV_2021)", "%", "S37", PCT),
    ("   subsidised-sale floor area (median 35 m² per household)", "=B26*B27*FLOOR_HH_PUBLIC", "m²", "S36", KWH0),
    ("   private floor area (median 47 m² per household)", "=B26*(1-B27)*FLOOR_HH_PRIVATE", "m²", "S36", KWH0),
    ("Method 2 - domestic floor area by housing type", "=B25+B28+B29", "m²", "Housing Department + Census medians", KWH0),
    ("   method 2 vs method 1", "=B30/B22-1", "%", "", PCT),
    ("District floor area per person (method 1)", "=B22/ST_POP", "m²/person", "", "0.0"),
]
for k, (lab, f, unit, basis, nf) in enumerate(fb):
    r = 18 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", f, "calc", nf, fill=FILL_KEY if r in (22, 30) else None, bold=r in (22, 30))
    put(ws, f"C{r}", unit, "text")
    put(ws, f"D{r}", basis, "text", wrap=True)
add_name("ST_AREA1", "District_Scaling!$B$22")
add_name("ST_AREA2", "District_Scaling!$B$30")

section(ws, 34, "C. Scaled-up residential electricity of Sha Tin District", 1, 7)
header(ws, 35, ["Method", "Factor", "Factor unit", "Scaling quantity", "Quantity unit", "District electricity (GWh/yr)", "vs EMSD benchmark"], height=45)
hwcook = "(HW_KWH+Inventory!J20)/OCC"
fc = [
    ("(1) Per capita: e_pc x population", "=DS_EPC", "kWh/person/yr", "=ST_POP", "persons", "=B36*D36/1e6"),
    ("(2) Per m²: e_m2 x floor area (method 1)", "=DS_EM2", "kWh/m²/yr", "=ST_AREA1", "m²", "=B37*D37/1e6"),
    ("(2b) Per m²: e_m2 x floor area (method 2)", "=DS_EM2", "kWh/m²/yr", "=ST_AREA2", "m²", "=B38*D38/1e6"),
    ("(3) Per household: household electricity x households", "=HH_KWH", "kWh/household/yr", "=ST_HH", "households", "=B39*D39/1e6"),
    ("(4) Composite: a x floor area + b x residents (method 1 area)", "=DS_A", "a, kWh/m²/yr", "=ST_AREA1", "m²", "=(DS_A*ST_AREA1+DS_B*ST_POP)/1e6"),
    ("(4b) Composite with method 2 floor area", "=DS_A", "a, kWh/m²/yr", "=ST_AREA2", "m²", "=(DS_A*ST_AREA2+DS_B*ST_POP)/1e6"),
    ("(4c) Composite, allowing for gas water heating and cooking", f"={hwcook}", "kWh/person/yr moved to gas", "=ST_POP*GAS_SHARE", "residents using gas", "=F40-B42*D42/1e6"),
    ("Benchmark A: HK residential electricity per person (EMSD 2021) x population", "=HK_RES_GWH*1e6/HK_POP_2021", "kWh/person/yr", "=ST_POP", "persons", "=B43*D43/1e6"),
    ("Benchmark B: CLP residential sales per account (2025) x households", "=CLP_RES_GWH*1e6/CLP_RES_ACCTS", "kWh/account/yr", "=ST_HH", "households", "=B44*D44/1e6"),
]
for k, (lab, f, unit, q, qu, gwh) in enumerate(fc):
    r = 36 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", f, "calc", KWH)
    put(ws, f"C{r}", unit, "text", wrap=True)
    put(ws, f"D{r}", q, "calc", KWH0)
    put(ws, f"E{r}", qu, "text")
    put(ws, f"F{r}", gwh, "calc", KWH, fill=FILL_KEY, bold=True)
    put(ws, f"G{r}", f"=F{r}/$F$43-1", "calc", PCT)
put(ws, "A46", "Per-m² estimate ÷ per-capita estimate", "text")
put(ws, "B46", "=F37/F36", "calc", "0.00")
put(ws, "C46", "equals district floor area per person ÷ audited floor area per person:", "note")
put(ws, "F46", "=B32/B9", "calc", "0.00")
put(ws, "A47", "Composite estimate ÷ per-capita estimate", "text")
put(ws, "B47", "=F40/F36", "calc", "0.00")
add_name("ST_GWH_PC", "District_Scaling!$F$36")
add_name("ST_GWH_M2", "District_Scaling!$F$37")
add_name("ST_GWH_COMP", "District_Scaling!$F$40")

section(ws, 50, "D. The composite factor: energy vs floor area and number of residents", 1, 7)
put(ws, "A51", "Household electricity E = a x A + b x N   (A = floor area, N = residents).  Per person: e(s) = b + a x s, where s = floor area per person.  Per m²: e(d) = a + b x d, where d = residents per m².", "note")
header(ws, 52, ["Household type", "Floor area per person s (m²)", "Residents N", "Floor area A (m²)", "Electricity E = aA + bN (kWh/yr)", "Per person (kWh/yr)", "Per m² (kWh/m²/yr)"], height=45)
fd = [
    ("Audited student room (this audit)", "=ROOM_AREA/OCC", "=OCC"),
    ("Sha Tin public rental flat", "=PRH_SPACE_PP", "=PRH_AHS"),
    ("Sha Tin average household (method 1)", "=B32", "=ST_AHS"),
    ("Private flat (47 m², Sha Tin household size)", "=FLOOR_HH_PRIVATE/ST_AHS", "=ST_AHS"),
]
for k, (lab, sval, n) in enumerate(fd):
    r = 53 + k
    put(ws, f"A{r}", lab, "text", wrap=True)
    put(ws, f"B{r}", sval, "calc", "0.0")
    put(ws, f"C{r}", n, "calc", "0.0")
    put(ws, f"D{r}", f"=B{r}*C{r}", "calc", "0.0")
    put(ws, f"E{r}", f"=DS_A*D{r}+DS_B*C{r}", "calc", KWH0)
    put(ws, f"F{r}", f"=E{r}/C{r}", "calc", KWH0)
    put(ws, f"G{r}", f"=E{r}/D{r}", "calc", KWH)
put(ws, "A58", "Check: the composite factor reproduces the audited household exactly (must be 0)", "note")
put(ws, "E58", "=ROUND(E53-HH_KWH,6)", "calc", "0.000")
put(ws, "A59", "Hong Kong average residential electricity per household, 2021 (EMSD ÷ Census households)", "text", wrap=True)
put(ws, "E59", "=HK_RES_GWH*1e6/HK_HH_2021", "calc", KWH0, fill=FILL_KEY)
put(ws, "A60", "Composite estimate for the Sha Tin average household vs the Hong Kong average", "text", wrap=True)
put(ws, "E60", "=E55/E59-1", "calc", PCT)
widths(ws, {"A": 58, "B": 14, "C": 16, "D": 16, "E": 16, "F": 16, "G": 14})
bar = BarChart()
bar.type = "bar"
bar.title = "Sha Tin residential electricity - scaling methods (GWh/yr)"
bar.add_data(Reference(ws, min_col=6, min_row=35, max_row=44), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=36, max_row=44))
bar.legend = None
bar.height, bar.width = 10, 18
ws.add_chart(bar, "I4")

# ======================================================================================
# EE-EAes SUMMARY (home adaptation)
ws = wb.create_sheet("EE-EAes_Summary")
title(ws, "Executive Summary - adapted from EMSD Technical Form EE-EAes (V.2)",
      "Residential buildings are not prescribed under Cap. 610; the form's structure is applied voluntarily to the audited home. Values link to the calculation sheets.")
r = 4
def form_row(label, value, kind="link", nf=None, note=None):
    global r
    put(ws, f"A{r}", label, "text", wrap=True)
    put(ws, f"B{r}", value, kind, nf, wrap=True)
    if note:
        put(ws, f"C{r}", note, "text", wrap=True)
        ws.merge_cells(f"C{r}:H{r}")
    elif kind == "text":
        ws.merge_cells(f"B{r}:H{r}")
    r += 1

section(ws, r, "Part 1 - (A) Administrative information (EAC Clause 8.1(m))", 1, 6); r += 1
form_row("Name of building (entity audited)", f"{AUDIT['building']} - {AUDIT['room']}", "text")
form_row("Address of building", AUDIT["district"], "text")
form_row("1) Date of commencement of energy audit", AUDIT["audit_start"], "text")
form_row("2) Date of completion of energy audit", AUDIT["audit_end"], "text")
form_row("3) Energy Audit Form validity", "N/A - voluntary home audit (not issued under Cap. 610)", "text")
form_row("4) Energy Audit Report reference no.", "SEE8115-HEA-2026-01", "text")
form_row("5) Energy imported from / exported to other buildings?", "No. The room is supplied from the hall's CLP supply; room AC is pre-paid by Octopus card.", "text")
r += 1
section(ws, r, "Part 1 - (B)(I) Building type, usage and operation", 1, 6); r += 1
form_row("1)(a) Type of building", "Others - residential student hostel (MiC, 2024)", "text")
form_row("1)(b) Portion of common area", "N/A - audited unit is a single room; shared washroom and kitchen allocated by users", "text")
form_row("1)(c) No. of blocks", "Village of 3 towers / 6 halls (S15); audited unit in Hall 13", "text")
form_row("2) Internal floor area of audited unit (m²)", "=ROOM_AREA", "link", "0.0")
form_row("3) Floor of audited unit", "10/F", "text")
form_row("4) Major type of facade", "MiC concrete modules; dark aluminium-framed single-glazed casement window facing NW (328°)", "text")
form_row("5) Occupation", "Opened Q3 2024; opening and naming ceremony Dec 2024 (S15)", "text")
form_row("6) Type of central air-conditioning", "Not applicable - unitary window AC (room) and split AC (kitchen)", "text")
form_row("7) Weekly occupied hours of the room", "=OCC_WKDAY_H*5+OCC_SAT_H+OCC_SUN_H", "link", "0")
form_row("   Daily average no. of occupants", "=OCC", "link", "0")
form_row("   Occupant density (m² per person)", "=ROOM_AREA/OCC", "link", "0.0")
r += 1
section(ws, r, "Part 1 - (B)(II) Building services installation of the audited unit", 1, 6); r += 1
header(ws, r, ["Equipment (type code)", "Cooling", "Compressor", "Refrigerant", "Rated capacity (kW)", "Input power (kW)", "Qty", "COP (kW/kW)"])
r += 1
put(ws, f"A{r}", "Window AC (R - room type)", "text")
put(ws, f"B{r}", "A (air-cooled)", "text")
put(ws, f"C{r}", "UA (inverter rotary)", "text")
put(ws, f"D{r}", "R32", "text")
put(ws, f"E{r}", "=AC_CAP_KW", "link", "0.00")
put(ws, f"F{r}", "=AC_RATED_KW", "link", "0.000")
put(ws, f"G{r}", 1, "in", "0")
put(ws, f"H{r}", f"=E{r}/F{r}", "calc", "0.00")
r += 1
form_row("Total lighting power of room (kW) (EAC Clause 8.1(f))", "=(CEIL_N*CEIL_W+DESK_N*DESK_W)/1000", "link", "0.000")
r += 1
section(ws, r, "Part 2 - Historical energy consumption analysis (EAC Clause 8.1(g))", 1, 6); r += 1
form_row("1) Annual electricity, past 1st 12-month (kWh/annum)", "=ROOM_KWH", "link", KWH,
         "Past 2nd and 3rd 12-month periods: N/A (room occupied from Jan 2026; EAC 8.1(g)(i))")
form_row("2) Annual energy other than electricity (MJ/annum)", 0, "in", KWH0, "No Towngas or LPG")
form_row("3) Annual total energy (MJ/annum)", "=ROOM_KWH*MJ_PER_KWH", "link", KWH0)
form_row("4) Annual EUI (MJ/m²/annum)", "=ROOM_EUI_MJ", "link", "0.0", "Equals kWh/m²/annum x 3.6")
put(ws, f"A{r}", "5) Monthly EUI (MJ/m²/month), typical year", "text", wrap=True)
r += 1
header(ws, r, MONTHS[:6], first_col=2)
r += 1
for i in range(6):
    put(ws, f"{chr(66 + i)}{r}", f"=Monthly!Q{6 + i}", "link", "0.00")
r += 1
header(ws, r, MONTHS[6:], first_col=2)
r += 1
for i in range(6):
    put(ws, f"{chr(66 + i)}{r}", f"=Monthly!Q{12 + i}", "link", "0.00")
r += 1
put(ws, f"A{r}", "6) Annual breakdown (MJ/annum)", "text", wrap=True)
r += 1
header(ws, r, ["Air-conditioning", "Lighting", "Lift & escalator", "Others"], first_col=2)
r += 1
put(ws, f"B{r}", "=Breakdown!F21", "link", KWH0)
put(ws, f"C{r}", "=Breakdown!F22", "link", KWH0)
put(ws, f"D{r}", "=Breakdown!F24", "link", KWH0)
put(ws, f"E{r}", "=Breakdown!F23", "link", KWH0)
r += 1
form_row("7) Energy supplied by shared building services (washroom, kitchen, laundry) as % of household total", "=SHARED_KWH/HH_KWH", "link", PCT)
form_row("8) Energy bill reference month", AUDIT["reference_month"] + " (Octopus pre-paid AC; no utility bill)", "text")
r += 1
section(ws, r, "Part 3 - Energy Management Opportunities (EMO)", 1, 6); r += 1
header(ws, r, ["Summary", "Category I", "Category II", "Category III"])
r += 1
put(ws, f"A{r}", "No. of EMOs", "text")
for k in range(3):
    put(ws, f"{chr(66 + k)}{r}", f"=EMO_Register!B{EMO_SUMMARY_ROW + 1 + k}", "link", "0")
r += 1
header(ws, r, ["Summary", "Lighting", "Air-conditioning", "Electrical", "Lift/Escalator", "Others"])
r += 1
put(ws, f"A{r}", "No. of EMOs", "text")
for k in range(5):
    put(ws, f"{chr(66 + k)}{r}", f"=EMO_Register!B{EMO_TYPE_ROW + 1 + k}", "link", "0")
r += 1
header(ws, r, ["Ref.", "Category", "Type", "Description of EMO (EAC Clause 8.1(l)(i))"])
r += 1
for k in range(len(emos)):
    er = first_emo + k
    put(ws, f"A{r}", f"=EMO_Register!A{er}", "link")
    put(ws, f"B{r}", f"=EMO_Register!B{er}", "link")
    put(ws, f"C{r}", f"=EMO_Register!C{er}", "link")
    put(ws, f"D{r}", f"=EMO_Register!E{er}", "link", wrap=True)
    ws.merge_cells(f"D{r}:H{r}")
    r += 1
widths(ws, {"A": 46, "B": 16, "C": 16, "D": 16, "E": 16, "F": 14, "G": 10, "H": 12})

# ======================================================================================
# EE-EAes FORM ENTRIES (what fill_eeaes_form.py writes into the official macro-enabled form; tick-boxes are
# given by their linked cell, or as 'Tick-box' for the three Part 1 boxes that have no linked cell)
ws = wb.create_sheet("EE-EAes_Form_Entries")
title(ws, "EE-EAes (V.2) form entries",
      "Entries made in 2_EE-EAes2018 (V.0)_SEE6115-8115_COMPLETED.xls (open it in Excel with macros enabled to see the form). "
      "Only sections relevant to a household are completed; CBSI and lift & escalator items are marked N/A as the form instructs.")
header(ws, 4, ["Form sheet", "Cell / control", "Field on the form", "Value entered", "Entry type", "Relevance / note"], height=30)
G = []
def g(sheet, cellref, field, val, kind, note="", nf=None):
    G.append((sheet, cellref, field, val, kind, note, nf))
NA_CBSI = "N/A - CBSI item (no central building services in a household)"
g("Part 1(A)", "C15", "Name of building", f"{AUDIT['room']}, {AUDIT['building']} (household energy audit)", "text", "Audited unit = the household's room")
g("Part 1(A)", "C16", "Address of building", "Lee Shau Kee Student Residence Village, Ma On Shan, Sha Tin District, New Territories, Hong Kong", "text")
g("Part 1(A)", "H17", "1) Date of commencement of energy audit", AUDIT["audit_start"], "date", "dd/mm/yyyy")
g("Part 1(A)", "H19", "2) Date of completion of energy audit", AUDIT["audit_end"], "date", "Within 6 months of the reference month (Sep 2026)")
g("Part 1(A)", "G21", "3) Energy Audit Form issued on", "N/A", "text", "Voluntary audit - no form issued under Cap. 610")
g("Part 1(A)", "J21", "3) Energy Audit Form expired on", "N/A", "text", "As above")
g("Part 1(A)", "H23", "4) Energy Audit Report reference no.", "SEE8115-HEA-2026-01", "text")
g("Part 1(A)", "Tick-box", "5) Import / export energy from/to other building?", "Tick 'No'", "tick", "Room supplied from the hall's CLP supply")
for c in ("E27", "H27", "J27"):
    g("Part 1(A)", c, "5) Building(s) importing / exporting energy", "N/A", "text", "Answer to item 5 is No")
g("Part 1(B)(I)", "-", "1)(a) Type of building", "Not ticked", "none", "Only commercial types are offered; the audited unit is residential (student hostel)")
g("Part 1(B)(I)", "H7", "1)(b) Portion of the building entity being common area (%)", 0, "number", "The audited room has no common area; shared facilities are allocated in the report", "0")
g("Part 1(B)(I)", "H8", "1)(c) No. of blocks of the building entity", 1, "number", "", "0")
g("Part 1(B)(I)", "H9", "2) Total internal floor area (m2)", "=ROOM_AREA", "number", "Room 1003B", "0.0")
g("Part 1(B)(I)", "H10", "3) No. of floors of the building entity", 1, "number", "Audited unit is on 10/F", "0")
g("Part 1(B)(I)", "Tick-box", "4) Major type of building facade", "Tick 'Non-curtain wall'", "tick", "MiC concrete modules with aluminium windows")
g("Part 1(B)(I)", "H12", "5) Date of issue of occupation approval", "Q3 2024 (residence opened; occupation-approval date not published)", "text")
g("Part 1(B)(I)", "Tick-box", "6) Type of central air-conditioning provided", "Tick 'Not applicable'", "tick", "Unitary window AC (room) and split AC (kitchen) only")
for rr in range(18, 24):
    g("Part 1(B)(I)", f"H{rr}", f"7) Average weekly operating hours - row {rr}", "N/A", "text", NA_CBSI)
    if rr != 22:
        g("Part 1(B)(I)", f"I{rr}", f"7) Daily average no. of occupants - row {rr}", "N/A", "text", NA_CBSI)
g("Part 1(B)(I)", "D60", "8)(f) Others (please specify)", "N/A - no CBSI-served area in the audited unit", "text", NA_CBSI)
g("Part 2", "E2", "1) Annual electricity, past 1st 12-month (kWh/annum)", "=ROUND(ROOM_KWH,1)", "number", "Room 1003B (typical 12 months to Sep 2026)", KWH)
g("Part 2", "G2", "1) Past 2nd 12-month", "N/A", "text", "Occupied from Jan 2026 (EAC 8.1(g)(i))")
g("Part 2", "I2", "1) Past 3rd 12-month", "N/A", "text", "Occupied from Jan 2026")
g("Part 2", "E5", "2) Other energy, past 1st 12-month (MJ/annum)", 0, "number", "No Towngas / LPG", "0")
g("Part 2", "G5", "2) Past 2nd 12-month", "N/A", "text", "")
g("Part 2", "I5", "2) Past 3rd 12-month", "N/A", "text", "")
ref_order = ["Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
slots = ["C14", "E14", "G14", "I14", "C16", "E16", "G16", "I16", "C18", "E18", "G18", "I18"]
for k, (slot, mon) in enumerate(zip(slots, ref_order)):
    g("Part 2", slot, f"5) Monthly EUI - month {k + 1} ({mon})", f"=ROUND(Monthly!Q{6 + MONTHS.index(mon)},2)", "number",
      "MJ/m2/month; month 12 = reference month (Sep 2026)", "0.00")
g("Part 2", "C20", "6) Breakdown - air-conditioning (MJ/annum)", "=ROUND(Breakdown!F21,0)", "number", "Room AC + desk fan", KWH0)
g("Part 2", "E20", "6) Breakdown - lighting (MJ/annum)", "=ROUND(Breakdown!F22,0)", "number", "", KWH0)
g("Part 2", "G20", "6) Breakdown - lift & escalator (MJ/annum)", "N/A", "text", "Lift & escalator item - not relevant to a household")
g("Part 2", "I20", "6) Breakdown - others (MJ/annum)", "=ROUND(Breakdown!F23,0)", "number", "Laptops, phone chargers", KWH0)
g("Part 2", "I22", "7) Energy supply from CBSI to building's units (%)", "N/A", "text", NA_CBSI)
g("Part 2", "I24", "8) Energy bill reference month", "30/09/2026", "date", "September 2026 (Octopus pre-paid AC; no utility bill)")
FORM_EMOS = ["EMO-01", "EMO-02", "EMO-04", "EMO-05", "EMO-06", "EMO-08", "EMO-10", "EMO-11", "EMO-12", "EMO-14", "EMO-18", "EMO-21"]
CAT_COL = {"I": "P", "II": "Q", "III": "R"}
TYPE_COL = {"LG": "S", "EL": "T", "AC": "U", "LE": "V", "Others": "W"}
emo_index = {e[0]: k for k, e in enumerate(emos)}
for k, ref in enumerate(FORM_EMOS):
    er = first_emo + emo_index[ref]
    _, cat, typ = emos[emo_index[ref]][:3]
    fr = 9 + 2 * k
    g("Part 3", f"A{fr}", f"EMO {chr(65 + k)} - reference", ref, "text")
    g("Part 3", f"{CAT_COL[cat]}{fr}", f"EMO {chr(65 + k)} - category", f"Tick Category {cat}", "tick")
    g("Part 3", f"{TYPE_COL[typ]}{fr}", f"EMO {chr(65 + k)} - type", f"Tick {typ}", "tick")
    tail = " Other EMOs (03, 07, 09, 13, 15, 16, 17, 19, 20) are in the report's EMO register (21 EMOs)." if k == len(FORM_EMOS) - 1 else ""
    desc = (f'=IF(EMO_Register!G{er}=0,EMO_Register!E{er}&". Viability statement only (EAC 7.4.4(d)); no saving claimed.",'
            f'EMO_Register!E{er}&". Saves "&TEXT(EMO_Register!G{er},"#,##0")&" kWh/yr"'
            f'&IF(ROUND(EMO_Register!G{er}-EMO_Register!H{er},1)<>0," (household share "&TEXT(EMO_Register!H{er},"#,##0")&")","")'
            f'&", HK$"&TEXT(EMO_Register!R{er},"#,##0")&"/yr, "&TEXT(EMO_Register!M{er},"#,##0")&" kg CO2e/yr; cost HK$"&TEXT(EMO_Register!J{er},"#,##0")'
            f'&"; payback "&IF(EMO_Register!J{er}=0,"immediate",TEXT(EMO_Register!K{er},"0.0")&" yr; NPV "&IF(EMO_Register!S{er}<0,"-","")&"HK$"&TEXT(ABS(EMO_Register!S{er}),"#,##0"))&".")'
            f'&"{tail}"')
    g("Part 3", f"H{fr}", f"EMO {chr(65 + k)} - description", desc, "text")
for k, (sheet, cellref, field, val, kind, note, nf) in enumerate(G):
    rr = 5 + k
    put(ws, f"A{rr}", sheet, "text")
    put(ws, f"B{rr}", cellref, "bold")
    put(ws, f"C{rr}", field, "text", wrap=True)
    put(ws, f"D{rr}", val, "link" if isinstance(val, str) and val.startswith("=") else "text", nf, wrap=True)
    put(ws, f"E{rr}", kind, "text")
    put(ws, f"F{rr}", note, "text", wrap=True)
widths(ws, {"A": 13, "B": 14, "C": 42, "D": 70, "E": 10, "F": 44})
ws.freeze_panes = "A5"

# ======================================================================================
# PHOTO LOG
ws = wb.create_sheet("Photo_Log")
title(ws, "Photo log - evidence for the inventory (files in photos/)")
header(ws, 4, ["Ref", "File", "What it shows", "Data extracted"])
PHOTOS = [
    ("01", "01_room_overview.jpg", "Room 1003B: twin beds, desks, ceiling LED battens, window AC above NW window", "Room layout; equipment present"),
    ("02a", "02a_window_view.jpg", "Window and outlook over Tolo Harbour / Ma On Shan", "Single glazing, dark aluminium frame, fly screen"),
    ("02b", "02b_window_fixed_pane_measured.jpg", "Bottom fixed pane measured 1.10 m x 0.17-0.19 m", "Glazing area 0.209 m²"),
    ("02c", "02c_window_casement_measured.jpg", "Casement measured 0.53 m x 1.26 m = 0.67 m²", "Two casements 1.34 m²"),
    ("02d", "02d_compass_328NW.jpg", "Compass 328° NW, 22°26'0\"N 114°14'32\"E, Ma On Shan, 44 m", "Orientation and location"),
    ("02e", "02e_curtain_closed.jpg", "Full-height grey curtain closed; AC at upper left", "Internal shading"),
    ("03a", "03a_ac_specification.jpg", "Midea MW-12CRF8E specification / energy label data", "3.34 kW, 1,192 W, R32, Grade 1, 362 kWh, CSPF 4.66"),
    ("03b", "03b_ac_front_louvres.jpg", "AC front: dark deposits on discharge louvres; Eco, Timer, Dry modes", "Maintenance finding; controls available"),
    ("05a", "05a_ceiling_led_battens.jpg", "Two linear LED ceiling battens", "Lighting inventory"),
    ("05b", "05b_desk_led_and_fan.jpg", "Under-shelf LED desk light and desk fan", "Task lighting"),
    ("06a", "06a_laptop_adapter_65W.jpg", "HP 65 W adapter TPN-CA17, Efficiency Level VI", "Laptop power"),
    ("06b", "06b_desk_fan_label_20W.jpg", "Fan label: ZW02-180, 180 mm, 20 W, 220 V", "Fan power"),
    ("08a", "08a_water_heater_DHB-E.jpg", "Stiebel Eltron DHB-E electronic instantaneous heater, 30-60 °C dial", "Hot-water system type"),
    ("08b", "08b_washroom_downlights.jpg", "Two LED downlights, exposed services, sensor", "Washroom lighting"),
    ("08c", "08c_washroom_light_and_fan.jpg", "Downlight and window exhaust fan", "Washroom lighting / ventilation"),
    ("08d", "08d_washroom_exhaust_fan.jpg", "J.E.E. window exhaust fan (runs 24 h)", "Ventilation finding"),
    ("08e", "08e_water_heater_spec.jpg", "Heater specification: 380 V, max 60 °C", "Electrical supply"),
    ("09a", "09a_dispenser_spec.jpg", "Aquatek AQ-1323T: 750 W heater, 160 W compressor, 10 L hot / 7.9 L cold / 3 L ice", "Dispenser ratings"),
    ("09b", "09b_water_dispenser.jpg", "Dispenser in kitchen, hot water at 98 °C; SRO cleaning notice", "Standing losses"),
    ("09c", "09c_induction_cooker_spec.jpg", "Philips HD4911: 2,100 W, EEL Grade 2", "Cooking power"),
    ("09d", "09d_induction_cooker_product.jpg", "Philips 5000 Series HD4911/80 (retailer image)", "Model identification"),
    ("09e", "09e_induction_energy_label.jpg", "Energy label Grade 2, thermal efficiency 88.1% (I220011/22)", "Cooking efficiency"),
    ("09f", "09f_kitchen_split_ac.jpg", "Toshiba split AC in kitchen (also used as photo 12 - left on 8 h+)", "Kitchen AC"),
    ("09g", "09g_kitchen_ac_energy_label.jpg", "Label: 688 kWh/yr @1,200 h, 7.00 kW, R410A, Grade 1", "Kitchen AC energy"),
    ("09h", "09h_microwave.jpg", "Toshiba MW2-MM24PC 24 L, 800 W output", "Kitchen inventory"),
    ("09i", "09i_fridge_energy_label.jpg", "Panasonic NR-BV361B: Grade 1, 263 kWh/yr, 233 + 99 L", "Refrigeration energy"),
]
for k, (ref, f, what, data) in enumerate(PHOTOS):
    rr = 5 + k
    put(ws, f"A{rr}", ref, "bold")
    put(ws, f"B{rr}", f, "text")
    put(ws, f"C{rr}", what, "text", wrap=True)
    put(ws, f"D{rr}", data, "text", wrap=True)
widths(ws, {"A": 6, "B": 36, "C": 60, "D": 42})

# ======================================================================================
# SOURCES
ws = wb.create_sheet("Sources")
title(ws, "Sources and references")
header(ws, 4, ["ID", "Short name", "Citation", "URL"])
for k, (sid, short, cit, url) in enumerate(SOURCES):
    rr = 5 + k
    put(ws, f"A{rr}", sid, "bold")
    put(ws, f"B{rr}", short, "text", wrap=True)
    put(ws, f"C{rr}", cit, "text", wrap=True)
    put(ws, f"D{rr}", url, "text", wrap=True)
    if url:
        ws[f"D{rr}"].hyperlink = url
widths(ws, {"A": 5, "B": 22, "C": 80, "D": 60})

# ======================================================================================
# README (front sheet)
ws = ws_readme
title(ws, "Home Energy Audit [2] - Dorm Room Energy & Carbon Audit Workbook",
      f"{AUDIT['auditor']} | {AUDIT['course']} | {AUDIT['building']}, {AUDIT['room']} | {AUDIT['district']}")
put(ws, "A4", "Purpose", "bold", border=False)
put(ws, "B4", "Systematic energy assessment of the auditor's home (a shared student-hostel room) following the EMSD Code of Practice for Building Energy Audit (EAC 2018) and the EPD/EMSD GHG Guidelines for Buildings in Hong Kong: baseline consumption, consumption patterns, carbon inventory and Energy Management Opportunities.", "text", wrap=True, border=False)
ws.merge_cells("B4:F4")
ws.row_dimensions[4].height = 45
put(ws, "A5", "Boundary", "bold", border=False)
put(ws, "B5", "Household = Room 1003B (15 m², 2 occupants) + allocated shares of the shared washroom (4 users) and floor kitchen (32 residents) + the household's own hall-laundry use. Hall common services (lifts, corridors, central plant) are outside the boundary.", "text", wrap=True, border=False)
ws.merge_cells("B5:F5")
ws.row_dimensions[5].height = 45
put(ws, "A6", "Period", "bold", border=False)
put(ws, "B6", "Data collected 1 Jan - 11 Oct 2026; results expressed as a typical 12-month baseline year.", "text", border=False)
ws.merge_cells("B6:F6")
put(ws, "A8", "Colour legend", "bold", border=False)
put(ws, "B8", "Blue text = input", "in")
put(ws, "C8", "Yellow fill = assumption", "assume")
put(ws, "D8", "Black = formula", "calc")
put(ws, "E8", "Green = link to another sheet", "link")
section(ws, 10, "Key results", 1, 6)
header(ws, 11, ["Indicator", "Value", "Unit", "Where calculated"])
keys = [
    ("Room AC energy (baseline, Octopus spend)", "=AC_KWH", "kWh/yr", "AC_Baseline", KWH0),
    ("Room AC energy, bottom-up energy-label check", "=AC_BU", "kWh/yr", "AC_Baseline", KWH0),
    ("Household electricity (room + allocated shared services)", "=HH_KWH", "kWh/yr", "Inventory", KWH0),
    ("  of which inside the 15 m² room", "=ROOM_KWH", "kWh/yr", "Inventory", KWH0),
    ("Electricity per person", "=HH_KWH/OCC", "kWh/person/yr", "Inventory", KWH0),
    ("Room EUI", "=ROOM_EUI_MJ", "MJ/m²/yr", "Monthly / Breakdown", "0.0"),
    ("Scope 1 (refrigerant screening)", "=S1_KG", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Scope 2 - CLP-specific factor 0.34", "=S2_CLP", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Scope 2 - territory default 0.7", "=S2_DEF", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Scope 3 (water, sewage, paper)", "=S3_KG", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Household total (CLP basis)", "=TOT_CLP", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Household total (default basis)", "=TOT_DEF", "kg CO2e/yr", "GHG_Inventory", KWH),
    ("Per person (CLP basis)", "=TOT_CLP/OCC", "kg CO2e/person/yr", "GHG_Inventory", KWH),
    ("Daily-activity extension (auditor)", "=LIFE_KG", "kg CO2e/yr", "Lifestyle_Scope3", KWH0),
    ("Auditor's personal total", "=PERSONAL_KG", "kg CO2e/yr", "Lifestyle_Scope3", KWH0),
    ("Category I EMO saving (household)", "=EMO_CAT1_KWH", "kWh/yr", "EMO_Register", KWH0),
    ("Category I EMO saving as % of household electricity", "=EMO_CAT1_KWH/HH_KWH", "%", "EMO_Register", PCT),
]
for k, (lab, f, unit, where, nf) in enumerate(keys):
    rr = 12 + k
    put(ws, f"A{rr}", lab, "text")
    put(ws, f"B{rr}", f, "calc", nf, fill=FILL_KEY, bold=True)
    put(ws, f"C{rr}", unit, "text")
    put(ws, f"D{rr}", where, "text")
rr = 12 + len(keys) + 1
section(ws, rr, "Sheet guide", 1, 6)
guide = [
    ("Inputs", "All inputs with unit, data quality and source; HKO 1991-2020 monthly temperatures"),
    ("AC_Baseline", "Room AC: top-down (Octopus spend) vs bottom-up (HK energy label) reconciliation"),
    ("HotWater", "Monthly shower hot-water energy from first principles"),
    ("Inventory", "Equipment inventory with allocation of shared services (EAC Step 1-2)"),
    ("Monthly", "Monthly profile, monthly EUI and emissions"),
    ("Breakdown", "End-use and EAC-category breakdown, KPIs and HK benchmarks"),
    ("GHG_Inventory", "Scope 1-2-3 inventory in HK GHG Guidelines table format"),
    ("Lifestyle_Scope3", "Auditor's commuting, flights and diet (reported separately)"),
    ("EMO_Register", "21 EMOs by EAC category: savings, cost, payback, NPV, SIR, cost of saved energy and CO2e"),
    ("Sensitivity", "One-at-a-time sensitivity of the key assumptions"),
    ("District_Scaling", "Assignment part 3: per-capita, per-m² and composite scale-up to Sha Tin District"),
    ("EE-EAes_Summary", "Executive summary in the structure of EMSD form EE-EAes"),
    ("EE-EAes_Form_Entries", "Values written into the official EE-EAes form (relevant sections only)"),
    ("Photo_Log", "Evidence photographs and the data taken from each"),
    ("Sources", "Full references with URLs"),
]
for k, (sh, d) in enumerate(guide):
    put(ws, f"A{rr + 1 + k}", sh, "bold")
    put(ws, f"B{rr + 1 + k}", d, "text")
    ws.merge_cells(f"B{rr + 1 + k}:F{rr + 1 + k}")
widths(ws, {"A": 50, "B": 18, "C": 20, "D": 20, "E": 26, "F": 10})

# Default Arial on any untouched cell, page setup
for sh in wb.worksheets:
    for rowc in sh.iter_rows():
        for c in rowc:
            if c.value is not None and (c.font is None or c.font.name != ARIAL):
                c.font = Font(name=ARIAL, size=10, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font else None)
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

wb.save(OUT)
print("saved", OUT, "names:", len(NAMES))

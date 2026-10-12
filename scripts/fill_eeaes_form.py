"""Fill a copy of the EMSD EE-EAes (V.2) executive-summary form with the audit results.

Usage: python3 fill_eeaes_form.py <original.xls> <workbook.xlsx> <output.xls>

The values come from the workbook's EE-EAes_Form_Entries sheet (recalculated). The form is a
macro-enabled BIFF8 workbook whose macros hide its sheets on save, so it is patched record by
record (xls_patch.py) rather than re-saved: the VBA project, form controls and very-hidden sheet
states stay exactly as issued by EMSD. Cached results of the form's own summary formulas are
updated so the values display correctly before Excel recalculates.
"""
import datetime
import sys

from openpyxl import load_workbook

from xls_patch import XlsPatcher, a1

# Part 1 tick-boxes have no linked cell, so they are found by their drawing-object id in the form
OBJ_TICKS = {
    ("Part 1(A)", "Tick 'No'"): 124,                    # 5) energy imported from / exported to other buildings
    ("Part 1(B)(I)", "Tick 'Non-curtain wall'"): 35,     # 4) major type of building facade
    ("Part 1(B)(I)", "Tick 'Not applicable'"): 12,       # 6) type of central air-conditioning
}

SRC, WB, OUT = sys.argv[1:4]

wb = load_workbook(WB, data_only=True)
ws = wb["EE-EAes_Form_Entries"]
p = XlsPatcher(SRC)
numbers = {}
part3_ticks = []
for r in range(5, ws.max_row + 1):
    sheet, ref, field, value, kind = (ws.cell(r, c).value for c in range(1, 6))
    if not sheet:
        continue
    if kind == "text":
        p.set_text(sheet, ref, str(value))
    elif kind == "number":
        p.set_number(sheet, ref, float(value))
        numbers[(sheet, ref)] = float(value)
    elif kind == "date":
        d = datetime.datetime.strptime(str(value), "%d/%m/%Y").date()
        p.set_date(sheet, ref, d)
    elif kind == "tick":
        if ref == "Tick-box":
            p.set_checkbox(sheet, OBJ_TICKS[(sheet, value)], True)
        else:  # tick-box linked to a cell (Part 3): tick it and set the linked cell TRUE
            p.set_checkbox(sheet, p.checkbox_for_link(sheet, ref), True)
            p.set_bool(sheet, ref, True)
            part3_ticks.append(a1(ref))
    elif kind != "none":
        raise ValueError(f"row {r}: unknown entry type {kind!r}")

# Part 2 summary formulas: total energy (E8) and EUI (E11) for the past 1st 12-month period
e8 = numbers[("Part 2", "E2")] * 3.6 + numbers[("Part 2", "E5")]
p.set_formula_result("Part 2", "E8", e8)
p.set_formula_result("Part 2", "E11", e8 / numbers[("Part 1(B)(I)", "H9")])

# Part 3 counters: P7:W7 = COUNTIF(column, TRUE); H2/I2/K2 categories, H4:L4 types (blank when 0)
counts = {col: 0 for col in range(15, 23)}  # P..W
for row, col in part3_ticks:
    counts[col] += 1
for col, n in counts.items():
    p.set_formula_result("Part 3", f"{chr(65 + col)}7", n)
for cell, col in (("H2", 15), ("I2", 16), ("K2", 17), ("H4", 18), ("I4", 20), ("J4", 19), ("K4", 21), ("L4", 22)):
    p.set_formula_result("Part 3", cell, counts[col] if counts[col] > 0 else "")

p.save(OUT)
print(f"wrote {OUT}: {len(p.log)} edits, {len(p.strings) - p.n_orig_strings} new strings, "
      f"categories I/II/III = {counts[15]}/{counts[16]}/{counts[17]}")
